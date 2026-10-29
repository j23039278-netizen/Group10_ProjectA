"""Checks for the gold annotation tools (run: python -m pytest tests -q).

Everything here uses a tiny synthetic sample, draft and skills list: no real
gold labels are read or written.
"""

import csv
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "data" / "jd_dataset" / "gold"))
import agreement  # noqa: E402
import collect_workbook  # noqa: E402
import make_workbook  # noqa: E402
from gold_common import Catalogue, evidence_in_text, load_skills, pick_b_refs  # noqa: E402

SAMPLE_FIELDS = ["jd_ref", "role_code", "title", "company", "location", "seniority", "source",
                 "source_url", "source_license", "collected_date", "raw_text"]
SKILL_FIELDS = ["skill_name", "category", "subcategory", "aliases", "ambiguous_aliases",
                "description", "source"]
LABEL_FIELDS = ["jd_ref", "skill_name", "importance", "evidence_text"]

SAMPLE = [
    {"jd_ref": "JD-DA-0001", "role_code": "DATA_ANALYST", "title": "Junior Data Analyst",
     "seniority": "junior",
     "raw_text": "Requirements:\n- Strong SQL and Python skills\n- Tableau is a plus\n=SUM(A1) is not a formula"},
    {"jd_ref": "JD-SD-0001", "role_code": "SOFTWARE_DEV", "title": "Graduate Software Engineer",
     "seniority": "entry", "raw_text": "We   use Docker daily.\n- Good  Communication skills"},
]
SKILLS = [
    ("Python", "Python 3", ""), ("SQL", "T-SQL", ""), ("Excel", "MS Excel", "excel"),
    ("Tableau", "", ""), ("Docker", "containers", ""), ("Communication", "", ""),
    ("PostgreSQL", "Postgres", ""), ("Git", "", ""), ("Linux", "", ""), ("AWS", "", ""),
]
DRAFT = [
    {"jd_ref": "JD-DA-0001", "skill_name": "SQL", "importance": "required",
     "evidence_text": "Strong SQL and Python skills"},
    {"jd_ref": "JD-DA-0001", "skill_name": "Python", "importance": "required",
     "evidence_text": "Strong SQL and Python skills"},
    {"jd_ref": "JD-DA-0001", "skill_name": "Excel", "importance": "required",
     "evidence_text": "Excel"},
    {"jd_ref": "JD-SD-0001", "skill_name": "Docker", "importance": "required",
     "evidence_text": "We use Docker daily"},
]


def write_rows(path, rows, fields):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields, restval="")
        w.writeheader()
        w.writerows(rows)
    return path


@pytest.fixture
def files(tmp_path):
    skills = [{"skill_name": n, "category": "Technical", "subcategory": "x", "aliases": al,
               "ambiguous_aliases": amb, "description": "d", "source": "v1"} for n, al, amb in SKILLS]
    return {
        "sample": write_rows(tmp_path / "sample.csv", SAMPLE, SAMPLE_FIELDS),
        "draft": write_rows(tmp_path / "draft.csv", DRAFT, LABEL_FIELDS),
        "skills": write_rows(tmp_path / "skills.csv", skills, SKILL_FIELDS),
        "dir": tmp_path,
    }


def catalogue(files):
    return Catalogue(load_skills(files["skills"]))


def header_row(ws):
    for row in ws.iter_rows(min_col=1, max_col=1):
        if row[0].value == "skill_name":
            return row[0].row
    raise AssertionError("no table header")


def texts():
    return {r["jd_ref"]: r["raw_text"] for r in SAMPLE}


# ── make_workbook ──────────────────────────────────────────────────────────
def test_workbook_a_is_prefilled_with_dropdowns(files):
    out = files["dir"] / "wb_a.xlsx"
    make_workbook.main(["--sample", str(files["sample"]), "--draft", str(files["draft"]),
                        "--skills", str(files["skills"]), "--out", str(out)])
    wb = load_workbook(out)
    assert wb.sheetnames == ["README", "JD-DA-0001", "JD-SD-0001", "skill_list"]
    assert wb["skill_list"].sheet_state == "hidden"
    assert "skill_names" in wb.defined_names

    ws = wb["JD-DA-0001"]
    h = header_row(ws)
    assert [c.value for c in ws[h]] == ["skill_name", "importance", "evidence_text", "keep", "note"]
    assert [ws.cell(row=h + i, column=1).value for i in (1, 2, 3)] == ["SQL", "Python", "Excel"]
    assert all(ws.cell(row=h + i, column=4).value is None for i in (1, 2, 3))  # keep left empty
    # the JD text is above the table, and a leading '=' stays text
    above = [ws.cell(row=r, column=1).value for r in range(1, h)]
    assert "- Strong SQL and Python skills" in above
    assert "=SUM(A1) is not a formula" in above
    assert all(c.data_type != "f" for row in ws.iter_rows(max_row=h) for c in row)

    rules = {str(dv.sqref)[0]: dv for dv in ws.data_validations.dataValidation}
    assert rules["A"].formula1 == "skill_names" and rules["A"].errorStyle == "warning"  # NEW: allowed
    assert rules["B"].formula1 == '"required,preferred,bonus"'
    assert rules["D"].formula1 == '"Y,N"'


def test_workbook_b_is_empty_and_never_overwritten(files):
    out = files["dir"] / "wb_b.xlsx"
    args = ["--annotator", "B", "--refs", "JD-SD-0001", "--sample", str(files["sample"]),
            "--draft", str(files["draft"]), "--skills", str(files["skills"]), "--out", str(out)]
    make_workbook.main(args)
    wb = load_workbook(out)
    assert wb.sheetnames == ["README", "JD-SD-0001", "skill_list"]
    ws = wb["JD-SD-0001"]
    h = header_row(ws)
    assert all(c.value is None for row in ws.iter_rows(min_row=h + 1) for c in row)
    with pytest.raises(SystemExit):
        make_workbook.main(args)


def test_pick_b_refs_two_per_role_fixed_seed():
    rows = [{"jd_ref": f"JD-{r}-{i:04d}", "role_code": r} for r in ("DA", "SD", "CY") for i in range(10)]
    refs = pick_b_refs(rows)
    assert len(refs) == 6 and refs == pick_b_refs(rows)
    assert sorted({ref.split("-")[1] for ref in refs}) == ["CY", "DA", "SD"]


# ── collect_workbook ───────────────────────────────────────────────────────
def test_round_trip_workbook_to_csv(files):
    wb_path = files["dir"] / "wb_a.xlsx"
    make_workbook.main(["--sample", str(files["sample"]), "--draft", str(files["draft"]),
                        "--skills", str(files["skills"]), "--out", str(wb_path)])
    wb = load_workbook(wb_path)
    ws = wb["JD-DA-0001"]
    h = header_row(ws)
    ws.cell(row=h + 1, column=4, value="Y")                    # SQL kept
    ws.cell(row=h + 2, column=4, value="y")                    # Python kept (lower case ok)
    ws.cell(row=h + 3, column=4, value="N")                    # Excel dropped
    for c, v in enumerate(["Tableau", "bonus", "Tableau is a plus", "Y"], start=1):
        ws.cell(row=h + 4, column=c, value=v)                  # added by hand
    ws = wb["JD-SD-0001"]
    h = header_row(ws)
    ws.cell(row=h + 1, column=4, value="Y")
    for c, v in enumerate(["communication", "required", "good communication skills", "Y"], start=1):
        ws.cell(row=h + 2, column=c, value=v)
    for c, v in enumerate(["NEW: CI/CD", "preferred", "we use docker", "Y"], start=1):
        ws.cell(row=h + 3, column=c, value=v)
    wb.save(wb_path)

    out = files["dir"] / "labels.csv"
    rc = collect_workbook.main([str(wb_path), "--sample", str(files["sample"]), "--draft",
                                str(files["draft"]), "--skills", str(files["skills"]), "--out", str(out)])
    assert rc == 0
    with open(out, newline="", encoding="utf-8") as f:
        rows = [(r["jd_ref"], r["skill_name"], r["importance"]) for r in csv.DictReader(f)]
    assert rows == [
        ("JD-DA-0001", "SQL", "required"), ("JD-DA-0001", "Python", "required"),
        ("JD-DA-0001", "Tableau", "bonus"), ("JD-SD-0001", "Docker", "required"),
        ("JD-SD-0001", "Communication", "required"), ("JD-SD-0001", "NEW: CI/CD", "preferred"),
    ]

    sheets = collect_workbook.read_workbook(wb_path)
    res = collect_workbook.collect(sheets, "A", texts(), catalogue(files))
    summary = {s["role"]: s for s in collect_workbook.role_summary(res, SAMPLE, DRAFT, sheets)}
    assert summary["DATA_ANALYST"] == {"role": "DATA_ANALYST", "jds_in_workbook": 1, "jds_labelled": 1,
                                       "rows": 3, "new_rows": 0, "dropped_from_draft": 1,
                                       "added_by_hand": 1}
    assert summary["SOFTWARE_DEV"]["new_rows"] == 1 and summary["SOFTWARE_DEV"]["added_by_hand"] == 2


def row(n, skill, importance="required", evidence="", keep="Y", note=""):
    return {"excel_row": n, "skill_name": skill, "importance": importance,
            "evidence_text": evidence, "keep": keep, "note": note}


def test_validation_errors_and_warnings(files):
    sheets = {
        "JD-DA-0001": [
            row(10, "SQL", evidence="strong   sql and PYTHON skills"),      # found (case/space)
            row(11, "Python", "preferred", "Python?"),                      # found, '?' ignored
            row(12, "Tableau", "bonus", "Tableau is required"),             # warning: not in text
            row(13, "Python", "bonus", "Strong SQL and Python skills"),     # duplicate -> merged
            row(14, "Excel", keep=""),                                      # A: not reviewed
            row(15, "Python 3", evidence="Python"),                         # error: alias
            row(16, "SQL", "must have", "SQL"),                             # error: importance
            row(17, "NEW: Postgres", evidence="SQL"),                       # error: in library
            row(18, "NEW: Data storytelling", "preferred", ""),             # ok, warning: no evidence
            row(19, "Git", keep="maybe"),                                   # error: keep value
        ],
        "JD-XX-9999": [row(10, "SQL")],                                     # error: unknown jd_ref
    }
    res = collect_workbook.collect(sheets, "A", texts(), catalogue(files))
    errors = "\n".join(f"{r} {n} {m}" for r, n, m in res.errors)
    warnings = "\n".join(f"{r} {n} {m}" for r, n, m in res.warnings)
    assert "15 'Python 3' is not a canonical name (did you mean 'Python'?)" in errors
    assert "16 importance must be one of" in errors
    assert "17 'NEW: Postgres' is already in the library as 'PostgreSQL'" in errors
    assert "19 keep must be Y or N" in errors
    assert "JD-XX-9999 None sheet name is not a jd_ref" in errors
    assert "12 evidence for Tableau not found" in warnings
    assert res.unreviewed == {"JD-DA-0001": 1}
    assert "18 no evidence_text" in warnings
    assert not any(n in (10, 11, 13) for _, n, _ in res.warnings)
    assert res.labels == []  # errors block the output

    res = collect_workbook.collect(sheets, "A", texts(), catalogue(files), skip_invalid=True)
    got = {r["skill_name"]: r["importance"] for r in res.labels}
    assert got == {"SQL": "required", "Python": "preferred", "Tableau": "bonus",
                   "NEW: Data storytelling": "preferred"}
    assert res.merged == 1


def test_annotator_b_keeps_all_non_empty_rows(files):
    sheets = {"JD-SD-0001": [row(5, "Docker", evidence="docker", keep=""),
                             row(6, "Linux", keep="N"),
                             row(7, "", "", "", keep="", note="unsure about k8s")]}
    res = collect_workbook.collect(sheets, "B", texts(), catalogue(files))
    assert [r["skill_name"] for r in res.labels] == ["Docker"] and not res.errors
    assert res.labelled == {"JD-SD-0001"}


def test_evidence_matching_ignores_case_whitespace_and_quotes():
    assert evidence_in_text("GOOD communication\n skills", "- Good  Communication skills")
    assert evidence_in_text("team’s tools?", "the team's tools")
    assert not evidence_in_text("", "anything")
    assert not evidence_in_text("Kubernetes", "We use Docker")


# ── agreement ──────────────────────────────────────────────────────────────
def test_cohen_kappa_hand_computed():
    # 2x2 table: both present 4, A only 1, B only 1, neither 4.
    # p_o = 8/10 = 0.8; p_e = 0.5*0.5 + 0.5*0.5 = 0.5; kappa = (0.8-0.5)/(1-0.5) = 0.6
    a = ["P"] * 5 + ["N"] * 5
    b = ["P"] * 4 + ["N", "P"] + ["N"] * 4
    assert agreement.cohen_kappa(a, b) == pytest.approx(0.6)
    assert agreement.cohen_kappa(a, a) == pytest.approx(1.0)
    assert agreement.cohen_kappa(["P", "P"], ["P", "P"]) != agreement.cohen_kappa(["P", "P"], ["P", "P"])  # nan


def test_weighted_kappa_hand_computed():
    # order bonus(0) < preferred(1) < required(2), weight |i-j|
    # observed = (0+1+0+0)/4 = 0.25; expected = 14/16 = 0.875; kappa = 1 - 0.25/0.875 = 5/7
    a = ["required", "required", "preferred", "bonus"]
    b = ["required", "preferred", "preferred", "bonus"]
    assert agreement.weighted_kappa(a, b, agreement.ORDINAL) == pytest.approx(5 / 7)


def test_agreement_units_and_outputs(tmp_path):
    a = [{"jd_ref": "JD-DA-0001", "skill_name": "SQL", "importance": "required", "evidence_text": "SQL"},
         {"jd_ref": "JD-DA-0001", "skill_name": "Python", "importance": "required", "evidence_text": "Python"},
         {"jd_ref": "JD-DA-0001", "skill_name": "NEW: Data storytelling", "importance": "bonus",
          "evidence_text": "stories"},
         {"jd_ref": "JD-SD-0001", "skill_name": "Docker", "importance": "required", "evidence_text": "Docker"}]
    b = [{"jd_ref": "JD-DA-0001", "skill_name": "SQL", "importance": "preferred", "evidence_text": "SQL"},
         {"jd_ref": "JD-DA-0001", "skill_name": "new: data  storytelling", "importance": "bonus",
          "evidence_text": "tell stories"},
         {"jd_ref": "JD-DA-0001", "skill_name": "Tableau", "importance": "bonus", "evidence_text": "Tableau"}]
    pa = write_rows(tmp_path / "a.csv", a, LABEL_FIELDS)
    pb = write_rows(tmp_path / "b.csv", b, LABEL_FIELDS)

    la, lb = agreement.load_labels(pa), agreement.load_labels(pb)
    names = ["Python", "SQL", "Tableau", "Docker", "Git", "Linux", "AWS", "Excel"]
    units = agreement.build_units(la, lb, ["JD-DA-0001"], names, neg_ratio=1.0, seed=42)
    union = [u for u in units if not u["negative"]]
    negs = [u for u in units if u["negative"]]
    assert len(union) == 4 and len(negs) == 4  # SQL, Python, NEW (same key), Tableau
    assert not {u["skill"] for u in negs} & {"SQL", "Python", "Tableau"}
    assert units == agreement.build_units(la, lb, ["JD-DA-0001"], names, neg_ratio=1.0, seed=42)

    s = agreement.presence_stats(union)
    assert (s["both"], s["a_only"], s["b_only"], s["neither"]) == (2, 1, 1, 0)
    assert s["percent"] == pytest.approx(50.0)
    imp = agreement.importance_stats(union)
    assert imp["units"] == 2 and imp["percent"] == pytest.approx(50.0)

    rc = agreement.main(["--a", str(pa), "--b", str(pb), "--out-dir", str(tmp_path)])
    assert rc == 0
    with open(tmp_path / "agreement_disagreements.csv", newline="", encoding="utf-8") as f:
        dis = {(r["skill"], r["A"], r["B"]) for r in csv.DictReader(f)}
    assert dis == {("SQL", "required", "preferred"), ("Python", "required", "absent"),
                   ("Tableau", "absent", "bonus")}  # JD-SD-0001 is A only, so not compared
    md = (tmp_path / "AGREEMENT.md").read_text(encoding="utf-8")
    assert "JDs labelled by both: **1**" in md and "weighted kappa" in md
