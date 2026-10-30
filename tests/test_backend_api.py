"""
Backend API tests (FastAPI TestClient) for SEAGAS.

Run from the repo root:
    python -m pip install -r tests/requirements.txt
    python -m pytest tests/test_backend_api.py -q

Each run creates its own database (seagas_test_api_<pid>) from database/schema.sql
and the seed files, and drops it at the end; seagas_db is never touched. The tests
are skipped when PostgreSQL is not reachable. Set TEST_PG_URL to point somewhere
else (default: postgresql://postgres:postgres123@localhost:5432/postgres).

Matching uses method="keyword" so results are exact and no NLP model is needed.
"""

import io
import os
import re
import sys
from pathlib import Path

import pytest

psycopg = pytest.importorskip("psycopg")
bcrypt = pytest.importorskip("bcrypt")
pytest.importorskip("fastapi")
pytest.importorskip("httpx")

ROOT = Path(__file__).resolve().parents[1]
PG_URL = os.environ.get("TEST_PG_URL", "postgresql://postgres:postgres123@localhost:5432/postgres")
DB_NAME = f"seagas_test_api_{os.getpid()}"
SQL_FILES = ["database/schema.sql", "database/migrate_existing_db.sql",
             "database/seed_skills_library_v2.sql", "database/seed_job_role_skills.sql",
             "database/seed_recommendations_library_final.sql", "database/seed_business_skills.sql"]
PASSWORD = "Test#Pass2026"


def run_sql_file(conn, path):
    # psql meta-commands (\set ...) are not SQL; the files are otherwise plain SQL
    sql = "\n".join(l for l in (ROOT / path).read_text(encoding="utf-8").splitlines()
                    if not l.lstrip().startswith("\\"))
    conn.execute(sql)


def db_url(name):
    return PG_URL.rsplit("/", 1)[0] + "/" + name


@pytest.fixture(scope="module")
def env():
    """A fresh database, the app wired to it, and a few accounts"""
    try:
        admin_conn = psycopg.connect(PG_URL, autocommit=True, connect_timeout=3)
    except Exception as e:
        pytest.skip(f"PostgreSQL not reachable ({e})")
    admin_conn.execute(f'DROP DATABASE IF EXISTS "{DB_NAME}"')
    admin_conn.execute(f'CREATE DATABASE "{DB_NAME}"')
    conn = psycopg.connect(db_url(DB_NAME), autocommit=True)
    for f in SQL_FILES:
        run_sql_file(conn, f)

    pw = bcrypt.hashpw(PASSWORD.encode(), bcrypt.gensalt()).decode()

    def user(email, role, name):
        return conn.execute("INSERT INTO users (email, password_hash, role, full_name) VALUES (%s,%s,%s,%s) "
                            "RETURNING user_id::text", (email, pw, role, name)).fetchone()[0]

    def student(email, name, skills):
        uid = user(email, "student", name)
        pid = conn.execute("INSERT INTO student_profiles (user_id, target_role) VALUES (%s,'Data Analyst') "
                           "RETURNING profile_id::text", (uid,)).fetchone()[0]
        for skill, level in skills:
            conn.execute("INSERT INTO student_skills (profile_id, skill_id, proficiency_level) "
                         "SELECT %s, skill_id, %s FROM skills_library WHERE skill_name = %s", (pid, level, skill))
        return uid, pid

    ids = {
        "alice": student("alice@student.newinti.edu.my", "Alice",
                         [("Python", "advanced"), ("SQL", "intermediate"), ("Excel", "beginner")]),
        "bob": student("bob@student.newinti.edu.my", "Bob", [("Java", "intermediate")]),
        "advisor": user("adv@advisor.newinti.edu.my", "advisor", "Advisor A"),
        "admin": user("boss@admin.newinti.edu.my", "admin", "Admin A"),
    }
    da = conn.execute("SELECT role_id::text FROM job_roles WHERE role_name='Data Analyst'").fetchone()[0]

    # the app reads these when it is imported
    os.environ["DATABASE_URL"] = db_url(DB_NAME).replace("postgresql://", "postgresql+psycopg://")
    os.environ["SECRET_KEY"] = "test-" + os.urandom(16).hex()
    sys.path.insert(0, str(ROOT / "backend"))
    from fastapi.testclient import TestClient
    import main
    import database

    client = TestClient(main.app)
    yield {"client": client, "conn": conn, "ids": ids, "da": da}

    client.close()
    database.engine.dispose()
    conn.close()
    admin_conn.execute(f'DROP DATABASE IF EXISTS "{DB_NAME}" WITH (FORCE)')
    admin_conn.close()


def student_number(env, email):
    row = env["conn"].execute("SELECT student_number FROM users WHERE lower(email) = lower(%s)", (email,)).fetchone()
    return row[0] if row else None


def login(env, email, password=PASSWORD, student_id=None):
    """POST /api/auth/login. Students sign in with their student ID too (looked up if not given)."""
    data = {"username": email, "password": password}
    sid = student_id if student_id is not None else student_number(env, email)
    if sid:
        data["student_id"] = sid
    return env["client"].post("/api/auth/login", data=data)


def token(env, email, password=PASSWORD, student_id=None):
    r = login(env, email, password, student_id)
    assert r.status_code == 200, r.text
    return {"Authorization": "Bearer " + r.json()["access_token"]}


@pytest.fixture(scope="module")
def alice(env):
    return token(env, "alice@student.newinti.edu.my")


@pytest.fixture(scope="module")
def bob(env):
    return token(env, "bob@student.newinti.edu.my")


@pytest.fixture(scope="module")
def advisor(env):
    return token(env, "adv@advisor.newinti.edu.my")


@pytest.fixture(scope="module")
def admin(env):
    return token(env, "boss@admin.newinti.edu.my")


@pytest.fixture(scope="module")
def biz_dean(env):
    """Dean of Business (sees Business students only). The default advisor is Dean of Computing."""
    pw = bcrypt.hashpw(PASSWORD.encode(), bcrypt.gensalt()).decode()
    env["conn"].execute("INSERT INTO users (email, password_hash, role, full_name, advisor_position, advisor_courses) "
                        "VALUES ('dean.biz@advisor.newinti.edu.my', %s, 'advisor', 'Dean Biz', 'dean', '{BUSINESS}') "
                        "ON CONFLICT (email) DO NOTHING", (pw,))
    return token(env, "dean.biz@advisor.newinti.edu.my")


def run(env, headers, **body):
    body.setdefault("matching_method", "keyword")
    return env["client"].post("/api/assessments/run", json=body, headers=headers)


# ── Authentication ───────────────────────────────────────────
class TestAuth:
    def test_login_returns_role_from_database(self, env):
        r = env["client"].post("/api/auth/login", data={"username": "ADV@advisor.newinti.edu.my", "password": PASSWORD})
        assert r.status_code == 200 and r.json()["role"] == "advisor"   # email is case-insensitive

    def test_wrong_password_and_unknown_email(self, env):
        c = env["client"]
        assert c.post("/api/auth/login", data={"username": "alice@student.newinti.edu.my", "password": "nope-nope"}).status_code == 401
        assert c.post("/api/auth/login", data={"username": "ghost@student.newinti.edu.my", "password": PASSWORD}).status_code == 401

    def test_overlong_password_is_401_not_500(self, env):
        r = env["client"].post("/api/auth/login", data={"username": "alice@student.newinti.edu.my", "password": "x" * 80})
        assert r.status_code == 401

    def test_student_login_needs_student_id(self, env):
        c = env["client"]
        sid = student_number(env, "alice@student.newinti.edu.my")
        assert len(sid) == 9 and sid.isdigit()                      # given by the database
        no_id = c.post("/api/auth/login", data={"username": "alice@student.newinti.edu.my", "password": PASSWORD})
        assert no_id.status_code == 401 and "student ID" in no_id.json()["detail"]
        assert login(env, "alice@student.newinti.edu.my", student_id="123456789").status_code == 401
        bobs = student_number(env, "bob@student.newinti.edu.my")
        assert login(env, "alice@student.newinti.edu.my", student_id=bobs).status_code == 401
        assert login(env, "alice@student.newinti.edu.my", student_id=f" {sid} ").status_code == 200
        reasons = {r[0] for r in env["conn"].execute(
            "SELECT details->>'reason' FROM audit_log WHERE action = 'auth.login_failed'").fetchall()}
        assert "wrong or missing student ID" in reasons

    def test_every_student_has_a_student_number(self, env):
        conn = env["conn"]
        assert conn.execute("SELECT count(*) FROM users WHERE role = 'student' AND student_number IS NULL").fetchone()[0] == 0
        assert conn.execute("SELECT count(*) FROM users WHERE role <> 'student' AND student_number IS NOT NULL").fetchone()[0] == 0
        with pytest.raises(psycopg.errors.CheckViolation):
            conn.execute("INSERT INTO users (email, password_hash, role, full_name, student_number) "
                         "VALUES ('bad@student.newinti.edu.my', 'x', 'student', 'B', '12AB')")
        synthetic = conn.execute("INSERT INTO users (email, password_hash, role, full_name) "
                                 "VALUES ('student0042@synthetic.example.com', 'x', 'student', 'S') "
                                 "RETURNING student_number").fetchone()[0]
        assert synthetic == "900000042"                             # same in every database

    @pytest.mark.parametrize("password", [
        "Sh0rt!",            # under 8 characters
        "nouppercase1!",     # no uppercase
        "NOLOWERCASE1!",     # no lowercase
        "NoDigitsHere!",     # no number
        "NoSymbols123",      # no symbol
        "Aa1!" + "x" * 70,   # over 72 bytes
        "密码Aa1!" * 13,      # over 72 bytes in UTF-8
    ])
    def test_new_password_rules(self, env, bob, password):
        r = env["client"].post("/api/auth/change-password", headers=bob,
                               json={"current_password": PASSWORD, "new_password": password})
        assert r.status_code == 400

    def test_voluntary_change_needs_current_password(self, env, bob):
        # Bob's password is not temporary, so he must prove he knows it
        r = env["client"].post("/api/auth/change-password", headers=bob,
                               json={"current_password": "wrong-one", "new_password": "Brand#New2026"})
        assert r.status_code == 400 and "current password" in r.json()["detail"]

    def test_self_registration_is_gone(self, env):
        r = env["client"].post("/api/auth/register", json={
            "email": "new1@student.newinti.edu.my", "password": "Password1!", "full_name": "N"})
        assert r.status_code in (404, 405)

    def test_deactivated_account_and_old_token(self, env, admin):
        c = env["client"]
        pw = bcrypt.hashpw(b"Password1!", bcrypt.gensalt()).decode()
        uid = env["conn"].execute("INSERT INTO users (email, password_hash, role, full_name) "
                                  "VALUES ('temp@student.newinti.edu.my', %s, 'student', 'T') RETURNING user_id::text",
                                  (pw,)).fetchone()[0]
        headers = token(env, "temp@student.newinti.edu.my", "Password1!")
        assert c.patch(f"/api/admin/users/{uid}", json={"is_active": False}, headers=admin).status_code == 200
        assert login(env, "temp@student.newinti.edu.my", "Password1!").status_code == 403
        assert c.get("/api/auth/me", headers=headers).status_code == 401


# ── Role-based access ────────────────────────────────────────
class TestRoles:
    @pytest.mark.parametrize("path", ["/api/admin/users", "/api/admin/audit-log", "/api/admin/skills",
                                      "/api/admin/resources"])
    def test_admin_api_needs_admin(self, env, alice, advisor, path):
        assert env["client"].get(path, headers=alice).status_code == 403
        assert env["client"].get(path, headers=advisor).status_code == 403

    def test_advisor_api_needs_staff(self, env, alice, advisor):
        assert env["client"].get("/api/advisor/cohort/overview", headers=alice).status_code == 403
        assert env["client"].get("/api/advisor/cohort/overview", headers=advisor).status_code == 200

    def test_no_token(self, env):
        assert env["client"].get("/api/students/profile").status_code == 401


# ── Assessments ──────────────────────────────────────────────
class TestAssessment:
    def test_role_template_assessment(self, env, alice):
        r = run(env, alice, role_id=env["da"])
        assert r.status_code == 201, r.text
        body = r.json()
        mr = body["match_results"]
        assert body["matching_method"] == "keyword" and len(mr) == body["skill_summary"]["total"] > 0
        w = {"required": 3, "preferred": 2, "bonus": 1}
        c = {"strong": 1, "developing": .5, "gap": 0}
        expected = round(sum(w[m["importance"]] * c[m["match_status"]] for m in mr)
                         / sum(w[m["importance"]] for m in mr) * 100, 2)
        assert body["readiness_score"] == expected
        sql = next(m for m in mr if m["skill_name"] == "SQL")
        assert sql["match_status"] == "strong" and sql["matched_student_skill"] == "SQL" and sql["reason"]
        excel = next((m for m in mr if m["skill_name"] == "Excel"), None)
        if excel:   # beginner with no evidence
            assert excel["match_status"] == "developing"

    def test_results_stored_with_reason(self, env, alice):
        aid = run(env, alice, role_id=env["da"]).json()["assessment_id"]
        n, with_reason = env["conn"].execute(
            "SELECT count(*), count(reason) FROM skill_match_results WHERE assessment_id=%s", (aid,)).fetchone()
        assert n > 0 and n == with_reason
        g = env["client"].get(f"/api/assessments/{aid}", headers=alice).json()
        assert all("reason" in m for m in g["match_results"])

    def test_history_has_dimension_scores(self, env, alice):
        run(env, alice, role_id=env["da"])
        h = env["client"].get("/api/assessments/my/history", headers=alice).json()[0]
        assert {"score_technical", "score_project_exp", "score_certification"} <= set(h)

    def test_unknown_method_is_400(self, env, alice):
        assert run(env, alice, role_id=env["da"], matching_method="spacy_ner").status_code == 400

    def test_experience_scores_follow_records(self, env, alice):
        c = env["client"]
        before = run(env, alice, role_id=env["da"]).json()["dimension_scores"]
        assert c.post("/api/students/projects", json={"project_title": "Sales dashboard"}, headers=alice).status_code == 201
        assert c.post("/api/students/certifications", json={"cert_name": "PL-300"}, headers=alice).status_code == 201
        after = run(env, alice, role_id=env["da"]).json()["dimension_scores"]
        assert after["project_experience"] > before["project_experience"]
        assert after["certification"] > before["certification"]

    def test_project_gives_evidence_without_changing_stored_skills(self, env):
        # The website has no evidence input for a skill; a project listing it counts (nlp/evidence.py)
        c, conn = env["client"], env["conn"]
        pw = bcrypt.hashpw(PASSWORD.encode(), bcrypt.gensalt()).decode()
        uid = conn.execute("INSERT INTO users (email, password_hash, role, full_name) VALUES "
                           "('dana@student.newinti.edu.my', %s, 'student', 'Dana') RETURNING user_id::text",
                           (pw,)).fetchone()[0]
        pid = conn.execute("INSERT INTO student_profiles (user_id, target_role) VALUES (%s, 'Data Analyst') "
                           "RETURNING profile_id::text", (uid,)).fetchone()[0]
        conn.execute("INSERT INTO student_skills (profile_id, skill_id, proficiency_level) "
                     "SELECT %s, skill_id, 'beginner' FROM skills_library WHERE skill_name = 'Power BI'", (pid,))
        dana = token(env, "dana@student.newinti.edu.my")
        power_bi = lambda: next(m for m in run(env, dana, role_id=env["da"]).json()["match_results"]
                                if m["skill_name"] == "Power BI")
        assert power_bi()["match_status"] == "developing"
        assert c.post("/api/students/projects", headers=dana, json={
            "project_title": "Sales Dashboard", "technologies": ["PowerBI"]}).status_code == 201
        m = power_bi()
        assert m["match_status"] == "strong"
        assert "beginner, with project evidence from Sales Dashboard" in m["reason"]
        stored = conn.execute("SELECT evidence_type FROM student_skills WHERE profile_id = %s", (pid,)).fetchone()[0]
        assert stored == []

    def test_pasted_jd_flow(self, env, alice):
        c = env["client"]
        jd = c.post("/api/jobs/descriptions", headers=alice, json={
            "title": "Junior Data Analyst",
            "raw_text": "Must know SQL and PowerBI; experience with data visualization tools is a plus. "
                        "Good communication skills."}).json()["jd_id"]
        a = c.post(f"/api/jobs/descriptions/{jd}/analyse?method=keyword", headers=alice)
        assert a.status_code == 200
        assert {"SQL", "Power BI", "Data Visualisation", "Communication"} <= {s["skill_name"] for s in a.json()["extracted_skills"]}
        again = c.post(f"/api/jobs/descriptions/{jd}/analyse?method=keyword", headers=alice).json()
        stored = env["conn"].execute("SELECT count(*) FROM jd_extracted_skills WHERE jd_id=%s", (jd,)).fetchone()[0]
        assert stored == again["skill_count"]           # re-analysis replaces, never duplicates
        r = run(env, alice, jd_id=jd)
        assert r.status_code == 201
        assert "PowerBI" not in {m["skill_name"] for m in r.json()["match_results"]}   # canonical names


# ── Record-level access ──────────────────────────────────────
class TestOwnership:
    def test_students_cannot_read_each_others_records(self, env, alice, bob, advisor):
        c = env["client"]
        aid = run(env, alice, role_id=env["da"]).json()["assessment_id"]
        assert c.get(f"/api/assessments/{aid}", headers=bob).status_code == 404
        assert c.get(f"/api/assessments/{aid}/recommendations", headers=bob).status_code == 404
        assert c.get(f"/api/assessments/{aid}", headers=advisor).status_code == 200
        jd = c.post("/api/jobs/descriptions", headers=alice, json={"raw_text": "We need a graduate with SQL and Python skills for reporting work."}).json()["jd_id"]
        assert c.get(f"/api/jobs/descriptions/{jd}", headers=bob).status_code == 404
        assert c.post(f"/api/jobs/descriptions/{jd}/analyse?method=keyword", headers=bob).status_code == 404
        assert run(env, bob, jd_id=jd).status_code == 404

    def test_malformed_id_is_404(self, env, alice):
        assert env["client"].get("/api/assessments/not-a-uuid", headers=alice).status_code == 404


# ── Student profile input ────────────────────────────────────
class TestProfile:
    @pytest.mark.parametrize("body", [{"gpa": 4.5}, {"gpa": -1}, {"year_of_study": 0}, {"year_of_study": 9}])
    def test_profile_limits(self, env, alice, body):
        assert env["client"].put("/api/students/profile", json=body, headers=alice).status_code == 422

    def test_skill_add_validation_and_delete(self, env, alice):
        c = env["client"]
        sid = env["conn"].execute("SELECT skill_id::text FROM skills_library WHERE skill_name='Tableau'").fetchone()[0]
        assert c.post("/api/students/skills", json={"skill_id": sid, "proficiency_level": "expert"}, headers=alice).status_code == 422
        assert c.post("/api/students/skills", json={"skill_id": "nope"}, headers=alice).status_code == 422
        assert c.post("/api/students/skills", json={"skill_id": sid, "proficiency_level": "beginner",
                                                    "evidence_type": ["project"]}, headers=alice).status_code == 201
        row = next(s for s in c.get("/api/students/skills", headers=alice).json() if s["skill_name"] == "Tableau")
        assert c.delete(f"/api/students/skills/{row['skill_id']}", headers=alice).status_code == 200
        assert "Tableau" not in {s["skill_name"] for s in c.get("/api/students/skills", headers=alice).json()}


# ── Admin ────────────────────────────────────────────────────
class TestAdmin:
    def test_create_advisor_emails_login_details(self, env, admin):
        c = env["client"]
        adv = {"full_name": "Dr Lim", "role": "advisor", "advisor_position": "hop", "advisor_courses": ["CS"],
               "personal_email": "lim.personal@example.com"}
        create = lambda b: c.post("/api/admin/users", headers=admin, json=b)
        assert create({**adv, "email": "x@gmail.com"}).status_code == 400                # not the advisor domain
        assert create({**adv, "personal_email": None}).status_code == 400                # nowhere to send it
        # Admin accounts still get the email and temporary password from the admin
        assert create({"full_name": "A", "role": "admin", "email": "a2@admin.newinti.edu.my"}).status_code == 400

        r = create(adv)
        assert r.status_code == 201, r.text
        u = r.json()
        # School email from the staff ID; the password is made by the system and only emailed
        assert u["email"] == f"{u['staff_number']}@advisor.newinti.edu.my" and u["email_status"] == "not_sent"
        assert "password" not in u
        mail = next(e for e in c.get("/api/intake/outbox", headers=admin).json()["emails"]
                    if e["to_address"] == "lim.personal@example.com")
        assert u["staff_number"] in mail["body"] and u["email"] in mail["body"]
        assert "Head of Programme, Computer Science" in mail["body"]
        temp = re.search(r"Temporary password:\s+(\S+)", mail["body"]).group(1)
        # First sign-in with the temporary password, then a new password is required
        r = login(env, u["email"], temp)
        assert r.status_code == 200 and r.json()["must_change_password"] is True
        h = {"Authorization": "Bearer " + r.json()["access_token"]}
        assert c.get("/api/advisor/students", headers=h).status_code == 403
        assert c.post("/api/auth/change-password", headers=h, json={"new_password": "Lim#Own2026"}).status_code == 200
        assert c.get("/api/advisor/students", headers=h).status_code == 200
        # The admin may choose another school email
        named = create({**adv, "email": "tan@advisor.newinti.edu.my", "personal_email": "tan.p@example.com"})
        assert named.status_code == 201 and named.json()["email"] == "tan@advisor.newinti.edu.my"

    def test_admin_changes_school_email(self, env, admin):
        c, conn = env["client"], env["conn"]
        pw = bcrypt.hashpw(PASSWORD.encode(), bcrypt.gensalt()).decode()
        uid = conn.execute("INSERT INTO users (email, password_hash, role, full_name) VALUES "
                           "('old.name@advisor.newinti.edu.my', %s, 'advisor', 'Renamed') RETURNING user_id::text",
                           (pw,)).fetchone()[0]
        patch = lambda e: c.patch(f"/api/admin/users/{uid}", headers=admin, json={"email": e})
        assert patch("new.name@gmail.com").status_code == 400                       # wrong domain for an advisor
        assert patch("new.name@student.newinti.edu.my").status_code == 400
        assert patch("adv@advisor.newinti.edu.my").status_code == 400               # taken
        r = patch("New.Name@advisor.newinti.edu.my")
        assert r.status_code == 200 and r.json()["email"] == "new.name@advisor.newinti.edu.my"
        assert login(env, "new.name@advisor.newinti.edu.my").status_code == 200    # signs in with the new email
        assert login(env, "old.name@advisor.newinti.edu.my").status_code == 401

    def test_seed_skill_locked_admin_skill_editable(self, env, admin):
        c = env["client"]
        excel = env["conn"].execute("SELECT skill_id::text FROM skills_library WHERE skill_name='Excel'").fetchone()[0]
        assert c.patch(f"/api/admin/skills/{excel}", json={"description": "x"}, headers=admin).status_code == 400
        new = c.post("/api/admin/skills", headers=admin, json={"skill_name": "Quantum Computing", "category": "AI_Digital",
                                                               "aliases": ["qc", "QC"]})
        assert new.status_code == 201 and new.json()["source"] == "admin" and new.json()["aliases"] == ["qc"]
        upd = c.patch(f"/api/admin/skills/{new.json()['skill_id']}", json={"description": "Qubits"}, headers=admin)
        assert upd.status_code == 200 and upd.json()["description"] == "Qubits"

    def test_admin_role_with_skills_can_be_assessed(self, env, admin, alice):
        c = env["client"]
        rid = c.post("/api/admin/roles", headers=admin, json={"role_name": "BI Developer"}).json()["role_id"]
        sids = [r[0] for r in env["conn"].execute(
            "SELECT skill_id::text FROM skills_library WHERE skill_name IN ('SQL','Power BI','Communication')").fetchall()]
        r = c.put(f"/api/admin/roles/{rid}/skills", headers=admin,
                  json=[{"skill_id": s, "importance": "required"} for s in sids])
        assert r.status_code == 200 and len(r.json()["skills"]) == 3
        assert run(env, alice, role_id=rid).json()["skill_summary"]["total"] == 3
        assert c.put(f"/api/admin/roles/{env['da']}/skills", headers=admin, json=[]).status_code == 400

    def test_audit_log_records_events(self, env, admin, advisor):
        c = env["client"]
        c.post("/api/auth/login", data={"username": "ghost@student.newinti.edu.my", "password": "whatever1"})
        alice_uid = env["ids"]["alice"][0]
        c.get(f"/api/advisor/students/{alice_uid}", headers=advisor)
        actions = {r[0] for r in env["conn"].execute("SELECT DISTINCT action FROM audit_log").fetchall()}
        assert {"auth.login", "auth.login_failed", "assessment.run", "jd.analyse",
                "advisor.view_student", "admin.user_create", "admin.user_update",
                "admin.skill_create", "admin.role_create"} <= actions
        leaked = env["conn"].execute("SELECT count(*) FROM audit_log WHERE details::text LIKE '%%$2b$%%'").fetchone()[0]
        assert leaked == 0
        page = c.get("/api/admin/audit-log?action=auth.login_failed", headers=admin).json()
        assert page["total"] >= 1 and all(e["action"] == "auth.login_failed" for e in page["entries"])


    def test_seed_resource_locked_admin_resource_editable(self, env, admin):
        c = env["client"]
        listed = c.get("/api/admin/resources", headers=admin)
        assert listed.status_code == 200 and len(listed.json()) > 0
        seed = next(r for r in listed.json() if r["source"] == "seed")
        assert c.patch(f"/api/admin/resources/{seed['resource_id']}", json={"cost": "$1"}, headers=admin).status_code == 400

        tableau = env["conn"].execute("SELECT skill_id::text FROM skills_library WHERE skill_name='Tableau'").fetchone()[0]
        body = {"skill_id": tableau, "resource_title": "Tableau Dashboards 101", "resource_type": "course",
                "provider": "Tableau", "url": "https://www.tableau.com/learn", "duration_hours": 6}
        new = c.post("/api/admin/resources", headers=admin, json=body)
        assert new.status_code == 201 and new.json()["source"] == "admin" and new.json()["skill_name"] == "Tableau"
        assert c.post("/api/admin/resources", headers=admin, json=body).status_code == 400   # same title, same skill
        rid = new.json()["resource_id"]
        upd = c.patch(f"/api/admin/resources/{rid}", headers=admin, json={"difficulty": "intermediate", "cost": "$20"})
        assert upd.status_code == 200 and upd.json()["difficulty"] == "intermediate" and upd.json()["cost"] == "$20"
        assert c.patch(f"/api/admin/resources/{rid}", headers=admin, json={"resource_type": None}).status_code == 400

        # Deactivating any resource (seed ones too) stops it being recommended
        off = c.patch(f"/api/admin/resources/{seed['resource_id']}", json={"is_active": False}, headers=admin)
        assert off.status_code == 200 and off.json()["is_active"] is False
        assert c.patch(f"/api/admin/resources/{seed['resource_id']}", json={"is_active": True}, headers=admin).status_code == 200
        actions = {r[0] for r in env["conn"].execute("SELECT DISTINCT action FROM audit_log").fetchall()}
        assert {"admin.resource_create", "admin.resource_update"} <= actions


# ── Advisor ──────────────────────────────────────────────────
class TestAdvisor:
    def test_profiles_complete_counts_complete_profiles_only(self, env, advisor, bob):
        c = env["client"]
        expected = lambda: env["conn"].execute(
            "SELECT count(*) FROM student_profiles sp JOIN users u USING (user_id) "
            "WHERE sp.is_complete AND u.role='student' AND u.is_active").fetchone()[0]
        complete = lambda: env["conn"].execute("SELECT is_complete FROM student_profiles WHERE profile_id = %s",
                                               (env["ids"]["bob"][1],)).fetchone()[0]
        # Bob has a skill but only a target role: not complete until the academic fields are filled in
        assert complete() is False
        before = c.get("/api/advisor/cohort/overview", headers=advisor).json()["profiles_complete"]
        assert before == expected()
        conn = env["conn"]
        conn.execute("UPDATE student_profiles SET programme = 'Bachelor of Computer Science (Software Engineering)' WHERE profile_id = %s",
                     (env["ids"]["bob"][1],))
        assert complete() is False                       # the student's own fields are still missing
        full = {"university": "Swinburne", "year_of_study": 2, "target_role": "Data Analyst"}
        assert c.put("/api/students/profile", json=full, headers=bob).status_code == 200
        assert complete() is True
        assert c.get("/api/advisor/cohort/overview", headers=advisor).json()["profiles_complete"] == before + 1
        # Without any skill the profile is incomplete again
        java = env["conn"].execute("SELECT skill_id::text FROM skills_library WHERE skill_name='Java'").fetchone()[0]
        assert c.delete(f"/api/students/skills/{java}", headers=bob).status_code == 200
        assert complete() is False
        assert c.post("/api/students/skills", json={"skill_id": java}, headers=bob).status_code == 201
        assert complete() is True

    def test_student_detail(self, env, advisor, alice):
        c = env["client"]
        run(env, alice, role_id=env["da"])
        r = c.get(f"/api/advisor/students/{env['ids']['alice'][0]}", headers=advisor)
        assert r.status_code == 200
        d = r.json()
        assert {"student", "assessment", "skills", "skill_gaps", "history", "projects",
                "certifications", "progress"} <= d.keys()
        assert d["assessment"]["role_name"] == "Data Analyst" and len(d["history"]) >= 1
        assert d["history"] == sorted(d["history"], key=lambda h: h["created_at"])   # oldest first, for the chart
        assert any(g.get("reason") for g in d["skill_gaps"])
        # Only students have a record here; malformed ids are 404 like the other ID routes
        assert c.get(f"/api/advisor/students/{env['ids']['advisor']}", headers=advisor).status_code == 404
        assert c.get("/api/advisor/students/not-a-uuid", headers=advisor).status_code == 404
        assert c.get(f"/api/advisor/students/{env['ids']['alice'][0]}", headers=alice).status_code == 403


# ── Student projects and certifications: edit and delete ─────
class TestProfileRecords:
    def test_project_edit_delete_and_ownership(self, env, bob, alice):
        c = env["client"]
        assert c.post("/api/students/projects", headers=bob, json={"project_title": "Bot", "project_type": "personal"}).status_code == 201
        pid = next(p for p in c.get("/api/students/projects", headers=bob).json() if p["project_title"] == "Bot")["project_id"]
        upd = {"project_title": "Chat Bot", "technologies": ["Python"], "project_type": "internship", "year_completed": 2026}
        assert c.put(f"/api/students/projects/{pid}", headers=alice, json=upd).status_code == 404   # not hers
        assert c.delete(f"/api/students/projects/{pid}", headers=alice).status_code == 404
        assert c.put(f"/api/students/projects/{pid}", headers=bob, json={"project_title": ""}).status_code == 422
        assert c.put(f"/api/students/projects/{pid}", headers=bob, json=upd).status_code == 200
        row = next(p for p in c.get("/api/students/projects", headers=bob).json() if p["project_id"] == pid)
        assert row["project_title"] == "Chat Bot" and row["technologies"] == ["Python"]
        counts = lambda: env["conn"].execute("SELECT project_count, internship_count FROM student_profiles "
                                             "WHERE profile_id = %s", (env["ids"]["bob"][1],)).fetchone()
        assert counts()[0] >= 1 and counts()[1] >= 1   # internship-type project counts as an internship
        n = counts()[0]
        assert c.delete(f"/api/students/projects/{pid}", headers=bob).status_code == 200
        assert counts()[0] == n - 1
        assert c.delete(f"/api/students/projects/{pid}", headers=bob).status_code == 404

    def test_certification_edit_delete_and_ownership(self, env, bob, alice):
        c = env["client"]
        assert c.post("/api/students/certifications", headers=bob, json={"cert_name": "AZ-900"}).status_code == 201
        cid = next(x for x in c.get("/api/students/certifications", headers=bob).json() if x["cert_name"] == "AZ-900")["cert_id"]
        upd = {"cert_name": "AZ-900 Azure Fundamentals", "issuer": "Microsoft", "issue_date": "2026-05-01"}
        assert c.put(f"/api/students/certifications/{cid}", headers=alice, json=upd).status_code == 404
        assert c.put(f"/api/students/certifications/{cid}", headers=bob, json=upd).status_code == 200
        row = next(x for x in c.get("/api/students/certifications", headers=bob).json() if x["cert_id"] == cid)
        assert row["issuer"] == "Microsoft" and row["issue_date"] == "2026-05-01"
        certs = lambda: env["conn"].execute("SELECT certification_count FROM student_profiles WHERE profile_id = %s",
                                            (env["ids"]["bob"][1],)).fetchone()[0]
        n = certs()
        assert c.delete(f"/api/students/certifications/{cid}", headers=alice).status_code == 404
        assert c.delete(f"/api/students/certifications/{cid}", headers=bob).status_code == 200
        assert certs() == n - 1


    def test_evidence_files(self, env, alice, bob, advisor):
        c = env["client"]
        png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 100
        pdf = b"%PDF-1.4\n" + b"x" * 100
        # The create response carries the new id, so the page can attach files straight away
        pid = c.post("/api/students/projects", headers=alice, json={"project_title": "Evidence demo"}).json()["project_id"]
        assert any(p["project_id"] == pid for p in c.get("/api/students/projects", headers=alice).json())
        up = lambda h, data, name="proof.png", url=f"/api/students/projects/{pid}/files": \
            c.post(url, headers=h, files={"file": (name, data, "image/png")})

        # Type comes from the content, not the name or the browser's Content-Type
        assert up(alice, b"<svg onload=alert(1)></svg>", "evil.png").status_code == 400
        assert up(alice, b"", "empty.png").status_code == 400
        assert up(alice, b"\x89PNG\r\n\x1a\n" + b"\x00" * (5 * 1024 * 1024)).status_code == 413
        first = up(alice, pdf, "..\\..\\cert")
        assert first.status_code == 201
        assert first.json()["content_type"] == "application/pdf" and first.json()["file_name"] == "cert.pdf"
        assert "data" not in first.json()
        assert up(bob, png).status_code == 404                                   # not Bob's project
        for _ in range(4):
            assert up(alice, png).status_code == 201
        assert up(alice, png).status_code == 400                                 # 6th file: limit is 5

        files = next(p for p in c.get("/api/students/projects", headers=alice).json() if p["project_id"] == pid)["files"]
        assert len(files) == 5 and all("data" not in f for f in files)
        fid = first.json()["file_id"]
        got = c.get(f"/api/students/files/{fid}", headers=alice)
        assert got.status_code == 200 and got.content == pdf
        assert got.headers["x-content-type-options"] == "nosniff" and got.headers["content-disposition"].startswith("attachment")
        assert c.get(f"/api/students/files/{fid}", headers=bob).status_code == 404
        assert c.get(f"/api/students/files/{fid}", headers=advisor).status_code == 200
        assert env["conn"].execute("SELECT count(*) FROM audit_log WHERE action='advisor.view_file' AND target_id=%s",
                                   (fid,)).fetchone()[0] == 1
        assert c.get("/api/students/files/not-a-uuid", headers=alice).status_code == 404
        detail = c.get(f"/api/advisor/students/{env['ids']['alice'][0]}", headers=advisor).json()
        assert len(next(p for p in detail["projects"] if p["project_id"] == pid)["files"]) == 5

        assert c.delete(f"/api/students/files/{fid}", headers=bob).status_code == 404
        assert c.delete(f"/api/students/files/{fid}", headers=alice).status_code == 200
        assert c.get(f"/api/students/files/{fid}", headers=alice).status_code == 404

        # Certifications take files too, and files go with their project when it is deleted
        cid = c.post("/api/students/certifications", headers=bob, json={"cert_name": "CCNA"}).json()["cert_id"]
        assert up(bob, png, url=f"/api/students/certifications/{cid}/files").status_code == 201
        assert len(next(x for x in c.get("/api/students/certifications", headers=bob).json()
                        if x["cert_id"] == cid)["files"]) == 1
        assert c.delete(f"/api/students/projects/{pid}", headers=alice).status_code == 200
        assert env["conn"].execute("SELECT count(*) FROM student_evidence_files WHERE project_id=%s",
                                   (pid,)).fetchone()[0] == 0


    def test_office_files_and_links(self, env, bob):
        import zipfile
        c = env["client"]

        def office(*parts):
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w") as z:
                for name in ("[Content_Types].xml",) + parts:
                    z.writestr(name, "<x/>")
            return buf.getvalue()

        pid = c.post("/api/students/projects", headers=bob, json={"project_title": "Office demo"}).json()["project_id"]
        up = lambda name, data: c.post(f"/api/students/projects/{pid}/files", headers=bob, files={"file": (name, data, "application/octet-stream")})
        docx = up("report", office("word/document.xml"))
        assert docx.status_code == 201 and docx.json()["file_name"] == "report.docx"
        assert docx.json()["content_type"].endswith("wordprocessingml.document")
        assert up("slides.pptx", office("ppt/presentation.xml")).json()["content_type"].endswith("presentationml.presentation")
        assert up("macro.docx", office("word/document.xml", "word/vbaProject.bin")).status_code == 400  # macros
        assert up("fake.docx", office("other.xml")).status_code == 400                                  # a zip, not Word
        assert up("clip.mp4", b"\x00\x00\x00\x18ftypmp42" + b"\x00" * 50).status_code == 400            # videos as links

        link = c.post(f"/api/students/projects/{pid}/links", headers=bob, json={"url": "https://github.com/bob/demo"})
        assert link.status_code == 201 and link.json()["file_name"] == "github.com/bob/demo"
        files = next(p for p in c.get("/api/students/projects", headers=bob).json() if p["project_id"] == pid)["files"]
        assert sorted(f["kind"] for f in files) == ["file", "file", "link"]


# ── Development progress ─────────────────────────────────────
class TestProgress:
    def test_give_up_unfinished_activity_only(self, env, alice, bob):
        c = env["client"]
        aid = run(env, alice, role_id=env["da"]).json()["assessment_id"]
        rid, rid2 = [r[0] for r in env["conn"].execute(
            "SELECT resource_id::text FROM recommendations_library ORDER BY resource_title LIMIT 2").fetchall()]

        assert c.post(f"/api/progress/{rid}/start?assessment_id={aid}", headers=alice).status_code == 201
        mine = c.get("/api/progress/my", headers=alice).json()["progress"]
        assert next(p for p in mine if p["resource_id"] == rid)["assessment_id"] == aid   # for "View in Recommendations"

        assert c.delete(f"/api/progress/{rid}", headers=bob).status_code == 404           # Bob never started it
        assert c.delete(f"/api/progress/{rid}", headers=alice).status_code == 200
        assert all(p["resource_id"] != rid for p in c.get("/api/progress/my", headers=alice).json()["progress"])
        assert c.delete(f"/api/progress/{rid}", headers=alice).status_code == 404         # already gone
        assert c.delete("/api/progress/not-a-uuid", headers=alice).status_code == 404

        # Completed activities stay on the record
        c.post(f"/api/progress/{rid2}/start?assessment_id={aid}", headers=alice)
        c.post(f"/api/progress/{rid2}/links", headers=alice, json={"url": "https://github.com/alice/ab-test"})
        assert c.post(f"/api/progress/{rid2}/complete", headers=alice, json={"status": "completed"}).status_code == 200
        assert c.delete(f"/api/progress/{rid2}", headers=alice).status_code == 404

    def test_complete_needs_evidence_and_advisor_sees_it(self, env, alice, bob, advisor):
        c = env["client"]
        aid = run(env, alice, role_id=env["da"]).json()["assessment_id"]
        rid = env["conn"].execute("SELECT resource_id::text FROM recommendations_library "
                                  "ORDER BY resource_title OFFSET 5 LIMIT 1").fetchone()[0]
        complete = lambda: c.post(f"/api/progress/{rid}/complete", headers=alice, json={"status": "completed"})
        link = lambda h, url, title=None: c.post(f"/api/progress/{rid}/links", headers=h, json={"url": url, "title": title})

        assert link(alice, "https://youtu.be/demo").status_code == 404           # not started yet
        assert complete().status_code == 400                                      # not started
        c.post(f"/api/progress/{rid}/start?assessment_id={aid}", headers=alice)
        r = complete()
        assert r.status_code == 400 and "evidence" in r.json()["detail"]          # no evidence yet

        for bad in ("javascript:alert(1)", "ftp://x.com/a", "www.no-scheme.com", "https://has space.com"):
            assert link(alice, bad).status_code == 400, bad
        assert link(bob, "https://youtu.be/demo").status_code == 404              # Bob never started it
        video = link(alice, "https://youtu.be/demo", "My demo video")
        assert video.status_code == 201
        assert video.json()["kind"] == "link" and video.json()["file_name"] == "My demo video"
        assert c.get(f"/api/students/files/{video.json()['file_id']}", headers=alice).status_code == 400  # a link, not a file
        assert complete().status_code == 200

        progress = c.get(f"/api/progress/assessment/{aid}", headers=alice).json()["progress"][rid]
        assert [e["link_url"] for e in progress["evidence"]] == ["https://youtu.be/demo"]
        # The only piece of evidence of a completed activity cannot be removed
        assert c.delete(f"/api/students/files/{video.json()['file_id']}", headers=alice).status_code == 400

        detail = c.get(f"/api/advisor/students/{env['ids']['alice'][0]}", headers=advisor).json()
        act = next(p for p in detail["progress"] if p["progress_id"] == progress["evidence"][0]["progress_id"])
        assert act["status"] == "completed" and act["files"][0]["link_url"] == "https://youtu.be/demo"


# ── Recommendations: cost tiers ──────────────────────────────
@pytest.mark.parametrize("cost,tier", [
    ("Free", "free"), (" free ", "free"), (None, "free"),
    ("Free to audit", "free_to_learn"), ("Free tier", "free_to_learn"), ("Free / $14/mo", "free_to_learn"),
    ("Free trial", "paid"), ("Paid ~$15", "paid"), ("$150 USD", "paid"), ("Low cost", "paid"),
])
def test_cost_tier(cost, tier):
    sys.path.insert(0, str(ROOT))
    from recommendation.engine import cost_tier
    assert cost_tier(cost) == tier


def test_recommendations_free_first(env, alice):
    aid = run(env, alice, role_id=env["da"]).json()["assessment_id"]
    recs = env["client"].get(f"/api/assessments/{aid}/recommendations", headers=alice).json()["recommendations"]
    order = {"free": 1, "free_to_learn": 2, "paid": 3}
    tiers = [[r["cost_tier"] for r in rec["resources"]] for rec in recs]
    assert any(tiers)
    assert all(t == sorted(t, key=order.get) for t in tiers)       # within each skill: free first


# ── New student intake: advisor list -> admin accounts -> first login ──
def upload(env, headers, name, content, mime="text/csv"):
    return env["client"].post("/api/intake/batches", headers=headers, files={"file": (name, content, mime)})


class TestIntake:
    CSV = ("Full Name,Phone Number,Personal Email,Course,Programme\n"
           "Mei Ling Tan,+60 12-345 6789,meiling.tan@example.com,Business,Bachelor of Business (Marketing)\n"
           "Arjun Nair,0123456780,arjun.nair@example.com,ict,bachelor of information and communications technology (web and mobile development)\n"
           "No Email,0123,not-an-email,CS,Bachelor of Computer Science (Software Engineering)\n"
           "Mei Ling Again,012,MEILING.TAN@example.com,Business,Bachelor of Business (Marketing)\n"
           "Siti Aminah,,siti.aminah@example.com,Computer Science,Bachelor of Computer Science (Software Engineering)\n"
           "No Course,,no.course@example.com,Medicine,Bachelor of Computer Science (Software Engineering)\n"
           "Wrong Programme,,wrong.prog@example.com,Business,Bachelor of Computer Science (Software Engineering)\n").encode()

    def test_full_intake_flow(self, env, advisor, admin, alice):
        c, conn = env["client"], env["conn"]

        # 1. The list is uploaded (it mixes courses, so by an admin); bad rows are skipped and explained
        assert upload(env, alice, "new.csv", self.CSV).status_code == 403          # students cannot
        r = upload(env, admin, "new.csv", self.CSV)
        assert r.status_code == 201, r.text
        assert r.json()["students"] == 3
        assert [s["line"] for s in r.json()["skipped"]] == [4, 5, 7, 8]            # bad email, duplicate, course, programme
        assert "course" in r.json()["skipped"][2]["error"]
        assert "Business programmes" in r.json()["skipped"][3]["error"]
        bid = r.json()["batch_id"]
        assert any(b["batch_id"] == bid and b["pending"] == 3 for b in c.get("/api/intake/batches", headers=admin).json())

        # 2. The admin lets the system suggest IDs and emails, edits one, rejects one
        assert c.post(f"/api/intake/batches/{bid}/assign", headers=advisor).status_code == 403
        rows = c.post(f"/api/intake/batches/{bid}/assign", headers=admin).json()["rows"]
        assert all(len(r["student_number"]) == 9 and r["school_email"] == f"{r['student_number']}@student.newinti.edu.my"
                   for r in rows)
        mei, arjun, siti = rows
        assert [r["course"] for r in rows] == ["BUSINESS", "ICT", "CS"]
        alice_id = student_number(env, "alice@student.newinti.edu.my")
        patch = lambda row, n, e: c.patch(f"/api/intake/rows/{row['row_id']}", headers=admin,
                                          json={"student_number": n, "school_email": e})
        assert patch(mei, "12345", "x@student.newinti.edu.my").status_code == 422           # not 9 digits
        assert patch(mei, "105000001", "mei@gmail.com").status_code == 400                  # wrong domain
        assert patch(mei, alice_id, "mei@student.newinti.edu.my").status_code == 400        # taken
        assert patch(mei, "105000001", "105000001@student.newinti.edu.my").status_code == 200
        assert c.post(f"/api/intake/rows/{siti['row_id']}/reject", headers=admin, json={"note": "duplicate"}).status_code == 200
        assert c.post(f"/api/intake/rows/{siti['row_id']}/reject", headers=admin, json={}).status_code == 404

        # 3. Accounts are created; each student's personal email gets the details
        done = c.post(f"/api/intake/batches/{bid}/create-accounts", headers=admin, json={}).json()
        assert len(done["created"]) == 2 and not done["failed"]
        courses = dict(env["conn"].execute("SELECT full_name, course FROM users WHERE full_name IN "
                                           "('Mei Ling Tan', 'Arjun Nair')").fetchall())
        assert courses == {"Mei Ling Tan": "BUSINESS", "Arjun Nair": "ICT"}
        assert {d["email_status"] for d in done["created"]} == {"not_sent"}               # no mail server in tests
        mail = next(e for e in c.get("/api/intake/outbox", headers=admin).json()["emails"]
                    if e["to_address"] == "meiling.tan@example.com")
        assert "105000001" in mail["body"] and "105000001@student.newinti.edu.my" in mail["body"]
        assert "Course:              Business" in mail["body"]
        temp = re.search(r"Temporary password:\s+(\S+)", mail["body"]).group(1)
        assert c.get("/api/intake/outbox", headers=advisor).status_code == 403

        # 4. First sign-in: email + student ID + temporary password, then a new password is required
        r = login(env, "105000001@student.newinti.edu.my", temp, student_id="105000001")
        assert r.status_code == 200 and r.json()["must_change_password"] is True
        h = {"Authorization": "Bearer " + r.json()["access_token"]}
        assert c.get("/api/students/profile", headers=h).status_code == 403              # blocked until changed
        assert c.get("/api/auth/me", headers=h).json()["student_number"] == "105000001"
        # The temporary password is not asked again: the student has just signed in with it
        assert re.search(r"[A-Z]", temp) and re.search(r"[a-z]", temp) and re.search(r"[0-9]", temp)             and re.search(r"[^A-Za-z0-9]", temp)                                       # meets the rules itself
        change = lambda new: c.post("/api/auth/change-password", headers=h, json={"new_password": new})
        assert change("MyOwnPass9").status_code == 400                                   # no symbol
        assert change(temp).status_code == 400                                           # same as temporary
        assert change("MyOwnPass9!").status_code == 200
        profile = c.get("/api/students/profile", headers=h)                               # allowed now
        assert profile.status_code == 200 and profile.json()["programme"] == "Bachelor of Business (Marketing)"
        assert login(env, "105000001@student.newinti.edu.my", temp, student_id="105000001").status_code == 401
        assert login(env, "105000001@student.newinti.edu.my", "MyOwnPass9!", student_id="105000001").status_code == 200

        # The same person cannot be put on a new list again
        again = upload(env, advisor, "again.csv",
                       b"Full Name,Personal Email,Course,Programme\nMei Ling,meiling.tan@example.com,CS,Bachelor of Computer Science (Software Engineering)\n")
        assert again.status_code == 400

        # 5. An admin reset makes the password temporary again
        uid = conn.execute("SELECT user_id::text FROM users WHERE student_number = '105000001'").fetchone()[0]
        assert c.post(f"/api/admin/users/{uid}/reset-password", headers=admin, json={"password": "Reset#2026"}).status_code == 200
        r = login(env, "105000001@student.newinti.edu.my", "Reset#2026", student_id="105000001")
        assert r.json()["must_change_password"] is True

        actions = {a[0] for a in conn.execute("SELECT DISTINCT action FROM audit_log").fetchall()}
        assert {"intake.upload", "intake.reject", "intake.accounts_create", "auth.password_change"} <= actions

    def test_xlsx_upload_and_bad_files(self, env, advisor):
        from openpyxl import Workbook
        wb = Workbook()
        wb.active.append(["Name", "Mobile", "Email", "Course", "Major"])
        wb.active.append(["Lim Wei", 60123456789, "lim.wei@example.com", "  Information and Communications Technology ",
                          "Bachelor of Information and Communications Technology (Web and Mobile Development)"])
        buf = io.BytesIO()
        wb.save(buf)
        r = upload(env, advisor, "students.xlsx", buf.getvalue(),
                   "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        assert r.status_code == 201 and r.json()["students"] == 1, r.text
        rows = env["client"].get(f"/api/intake/batches/{r.json()['batch_id']}", headers=advisor).json()["rows"]
        assert rows[0]["phone"] == "60123456789"                                  # Excel number for the phone
        assert rows[0]["course"] == "ICT" and rows[0]["programme"] == "Bachelor of Information and Communications Technology (Web and Mobile Development)"
        assert upload(env, advisor, "x.xlsx", b"not a zip").status_code == 400
        assert upload(env, advisor, "x.txt", b"a,b").status_code == 400
        assert upload(env, advisor, "x.csv", b"Phone,Address\n1,2\n").status_code == 400  # no name / email column
        no_course = upload(env, advisor, "x.csv", b"Full Name,Personal Email\nA B,a.b@example.com\n")
        assert no_course.status_code == 400 and "Course" in no_course.json()["detail"]
        no_prog = upload(env, advisor, "x.csv", b"Full Name,Personal Email,Course\nA B,a.b@example.com,CS\n")
        assert no_prog.status_code == 400 and "Programme" in no_prog.json()["detail"]


# ── Courses (Computer Science / Business / ICT) ──────────────
class TestCourses:
    def test_every_student_has_a_course(self, env):
        conn = env["conn"]
        assert conn.execute("SELECT count(*) FROM users WHERE role = 'student' AND course IS NULL").fetchone()[0] == 0
        assert conn.execute("SELECT count(*) FROM users WHERE role <> 'student' AND course IS NOT NULL").fetchone()[0] == 0
        # Synthetic students are split by number: N mod 3 -> 1 CS, 2 Business, 0 ICT
        got = {n: conn.execute("INSERT INTO users (email, password_hash, role, full_name) "
                               "VALUES (%s, 'x', 'student', 'S') RETURNING course",
                               (f"student{n:04d}@synthetic.example.com",)).fetchone()[0]
               for n in (4000, 4001, 4002)}
        assert got == {4000: "CS", 4001: "BUSINESS", 4002: "ICT"}   # 4000 mod 3 = 1, 4001 -> 2, 4002 -> 0
        with pytest.raises(psycopg.errors.CheckViolation):
            conn.execute("UPDATE users SET course = 'MEDICINE' WHERE email = 'student4000@synthetic.example.com'")

    def test_course_in_lists_and_filters(self, env, advisor, biz_dean, admin, alice):
        c, conn = env["client"], env["conn"]
        conn.execute("UPDATE users SET course = 'BUSINESS' WHERE email = 'bob@student.newinti.edu.my'")
        names = lambda h: {s["full_name"]: s["course"] for s in c.get("/api/advisor/students", headers=h).json()
                           if s["full_name"] in ("Alice", "Bob")}
        assert names(advisor) == {"Alice": "CS"}             # Dean of Computing
        assert names(biz_dean) == {"Bob": "BUSINESS"}        # Dean of Business
        biz = c.get("/api/admin/users?course=business", headers=admin).json()
        assert biz["total"] >= 1 and all(u["course"] == "BUSINESS" for u in biz["users"])
        assert "Bob" in {u["full_name"] for u in biz["users"]}
        assert c.get("/api/admin/stats", headers=admin).json()["courses"]["BUSINESS"] >= 1
        assert c.get("/api/auth/me", headers=alice).json()["course"] == "CS"
        detail = c.get(f"/api/advisor/students/{env['ids']['alice'][0]}", headers=advisor).json()
        assert detail["student"]["course"] == "CS"


def course_student(env, email, name, course):
    """A student of one course, signed in"""
    pw = bcrypt.hashpw(PASSWORD.encode(), bcrypt.gensalt()).decode()
    env["conn"].execute("INSERT INTO users (email, password_hash, role, full_name, course) "
                        "VALUES (%s, %s, 'student', %s, %s) ON CONFLICT (email) DO NOTHING", (email, pw, name, course))
    return token(env, email)


class TestRolesByCourse:
    def test_roles_offered_per_course(self, env):
        c = env["client"]
        names = lambda course: {r["role_name"] for r in c.get(f"/api/jobs/roles?course={course}").json()}
        it = {"Web Developer", "Data Analyst", "Cloud Engineer", "Cybersecurity Analyst"}
        biz = {"Business Analyst", "Marketing Executive", "Financial Analyst", "HR Executive"}
        assert it | {"Software Developer", "Data Scientist"} <= names("CS")
        assert names("ICT") == it
        assert names("BUSINESS") == biz
        # Business role templates (seed_job_role_skills.sql v3): 9 required + 8 preferred + 6 bonus
        assert all(r["skill_count"] == 23 for r in c.get("/api/jobs/roles?course=BUSINESS").json())

    def test_business_student_assessed_against_business_role(self, env):
        c, conn = env["client"], env["conn"]
        farah = course_student(env, "farah@student.newinti.edu.my", "Farah", "BUSINESS")
        sid = lambda name: conn.execute("SELECT skill_id::text FROM skills_library WHERE skill_name = %s",
                                        (name,)).fetchone()[0]
        for skill, level in (("Recruitment", "advanced"), ("Payroll", "beginner")):
            assert c.post("/api/students/skills", headers=farah,
                          json={"skill_id": sid(skill), "proficiency_level": level}).status_code == 201
        assert c.put("/api/students/profile", headers=farah, json={"target_role": "HR Executive"}).status_code == 200
        hr = conn.execute("SELECT role_id::text FROM job_roles WHERE role_code = 'HR_EXEC'").fetchone()[0]
        r = run(env, farah, role_id=hr)
        assert r.status_code == 201, r.text
        status = {m["skill_name"]: m["match_status"] for m in r.json()["match_results"]}
        assert len(status) == 23
        assert (status["Recruitment"], status["Payroll"], status["Employment Law"]) == ("strong", "developing", "gap")

    def test_business_roles_are_seed_managed(self, env, admin):
        """Their skills come from seed_job_role_skills.sql, so admins can only switch them on or off"""
        c, conn = env["client"], env["conn"]
        rows = conn.execute("SELECT role_id::text, source FROM job_roles WHERE role_code IN "
                            "('BUSINESS_ANALYST', 'MARKETING_EXEC', 'FINANCIAL_ANALYST', 'HR_EXEC')").fetchall()
        assert len(rows) == 4 and {s for _, s in rows} == {"seed"}
        rid = rows[0][0]
        assert c.put(f"/api/admin/roles/{rid}/skills", headers=admin, json=[]).status_code == 400
        assert c.patch(f"/api/admin/roles/{rid}", headers=admin, json={"description": "x"}).status_code == 400
        assert c.patch(f"/api/admin/roles/{rid}", headers=admin, json={"is_active": False}).status_code == 200
        assert c.patch(f"/api/admin/roles/{rid}", headers=admin, json={"is_active": True}).status_code == 200

    def test_business_recommendations(self, env):
        """Business gaps get Business learning resources (seed_business_skills.sql)"""
        c, conn = env["client"], env["conn"]
        h = course_student(env, "fiona@student.newinti.edu.my", "Fiona", "BUSINESS")
        assert c.put("/api/students/profile", headers=h, json={"target_role": "Financial Analyst"}).status_code == 200
        for skill in ("Excel", "Accounting"):
            sid = conn.execute("SELECT skill_id::text FROM skills_library WHERE skill_name = %s", (skill,)).fetchone()[0]
            assert c.post("/api/students/skills", headers=h, json={"skill_id": sid}).status_code == 201
        fa = conn.execute("SELECT role_id::text FROM job_roles WHERE role_code = 'FINANCIAL_ANALYST'").fetchone()[0]
        r = run(env, h, role_id=fa)
        assert r.status_code == 201, r.text
        result = r.json()
        names = {m["skill_name"]: m["match_status"] for m in result["match_results"]}
        assert names["Excel"] != "gap" and names["Financial Modelling"] == "gap"
        recs = c.get(f"/api/assessments/{result['assessment_id']}/recommendations", headers=h).json()
        skills = {rec["skill_name"] for rec in recs["recommendations"]}
        assert {"Financial Modelling", "Financial Reporting", "Budgeting & Forecasting"} <= skills
        # Every Business skill has at least one active learning resource
        assert conn.execute("SELECT count(*) FROM skills_library sl WHERE source = 'v3' AND NOT EXISTS "
                            "(SELECT 1 FROM recommendations_library r WHERE r.skill_id = sl.skill_id AND r.is_active)"
                            ).fetchone()[0] == 0

    def test_target_role_must_suit_the_course(self, env):
        c = env["client"]
        carol = course_student(env, "carol@student.newinti.edu.my", "Carol", "BUSINESS")
        dan = course_student(env, "dan@student.newinti.edu.my", "Dan", "ICT")
        put = lambda h, role: c.put("/api/students/profile", headers=h, json={"target_role": role})
        assert put(carol, "Data Analyst").status_code == 400
        assert put(carol, "Business Analyst").status_code == 200
        assert put(dan, "Data Scientist").status_code == 400
        assert put(dan, "Data Analyst").status_code == 200
        assert put(dan, "Pilot").status_code == 400
        # An assessment against a role of another course is refused
        r = run(env, carol, role_id=env["da"])
        assert r.status_code == 400 and "course" in r.json()["detail"]

    def test_admin_role_courses(self, env, admin):
        c = env["client"]
        body = {"role_name": "Retail Manager", "courses": ["BUSINESS"]}
        new = c.post("/api/admin/roles", headers=admin, json=body)
        assert new.status_code == 201 and new.json()["courses"] == ["BUSINESS"]
        rid = new.json()["role_id"]
        assert "Retail Manager" in {r["role_name"] for r in c.get("/api/jobs/roles?course=BUSINESS").json()}
        upd = c.patch(f"/api/admin/roles/{rid}", headers=admin, json={"courses": ["ICT", "BUSINESS", "ICT"]})
        assert upd.status_code == 200 and upd.json()["courses"] == ["BUSINESS", "ICT"]
        assert c.patch(f"/api/admin/roles/{rid}", headers=admin, json={"courses": []}).status_code == 422
        assert c.patch(f"/api/admin/roles/{rid}", headers=admin, json={"courses": ["LAW"]}).status_code == 422
        # Without courses a new role is for Computer Science students
        cs = c.post("/api/admin/roles", headers=admin, json={"role_name": "Game Developer"})
        assert cs.json()["courses"] == ["CS"]


class TestResultsReminder:
    def test_semester_start(self, env):
        from datetime import date
        from routers.students import current_semester_start
        assert current_semester_start(date(2026, 3, 15)) == date(2026, 1, 1)
        assert current_semester_start(date(2026, 5, 1)) == date(2026, 5, 1)
        assert current_semester_start(date(2026, 8, 31)) == date(2026, 5, 1)
        assert current_semester_start(date(2026, 12, 31)) == date(2026, 9, 1)

    def test_gpa_cgpa_and_reminder(self, env):
        c, conn = env["client"], env["conn"]
        h = course_student(env, "erin@student.newinti.edu.my", "Erin", "CS")
        assert c.post("/api/students/profile", headers=h, json={"target_role": "Data Scientist"}).status_code in (200, 201)
        profile = lambda: c.get("/api/students/profile", headers=h).json()
        assert profile()["results_reminder"] is True                       # results never entered
        assert c.put("/api/students/profile", headers=h, json={"cgpa": 4.5}).status_code == 422
        body = {"target_role": "Data Scientist", "gpa": 3.6, "cgpa": 3.42}
        assert c.put("/api/students/profile", headers=h, json=body).status_code == 200
        p = profile()
        assert float(p["gpa"]) == 3.6 and float(p["cgpa"]) == 3.42 and p["results_reminder"] is False

        # A new semester starts: the reminder comes back until the student answers it
        conn.execute("UPDATE student_profiles sp SET results_updated_at = NOW() - INTERVAL '200 days' "
                     "FROM users u WHERE u.user_id = sp.user_id AND u.email = 'erin@student.newinti.edu.my'")
        assert profile()["results_reminder"] is True
        # Saving other fields with the same results does not count as an update
        assert c.put("/api/students/profile", headers=h, json={**body, "university": "INTI"}).status_code == 200
        assert profile()["results_reminder"] is True
        assert c.post("/api/students/profile/results-confirmed", headers=h).status_code == 200
        assert profile()["results_reminder"] is False


class TestCourseOverview:
    def test_overview_and_gaps_by_course(self, env, admin, advisor):
        c, conn = env["client"], env["conn"]
        assert c.get("/api/advisor/cohort/overview?course=BUSINESS", headers=advisor).status_code == 403
        assert c.get("/api/advisor/cohort/overview", headers=advisor).json()["courses"] == ["CS", "ICT"]
        advisor = admin                                       # admins see all three courses
        for course in ("CS", "BUSINESS", "ICT"):
            o = c.get(f"/api/advisor/cohort/overview?course={course.lower()}", headers=advisor)
            assert o.status_code == 200 and o.json()["course"] == course
            expected = conn.execute("SELECT count(*) FROM users WHERE role = 'student' AND is_active "
                                    "AND course = %s", (course,)).fetchone()[0]
            assert o.json()["total_students"] == expected
            assert c.get(f"/api/advisor/cohort/skill-gaps?course={course}", headers=advisor).status_code == 200
        everyone = c.get("/api/advisor/cohort/overview", headers=advisor).json()
        assert everyone["course"] is None
        assert everyone["total_students"] == sum(
            c.get(f"/api/advisor/cohort/overview?course={k}", headers=advisor).json()["total_students"]
            for k in ("CS", "BUSINESS", "ICT"))
        assert c.get("/api/advisor/cohort/overview?course=LAW", headers=advisor).status_code == 400
        assert c.get("/api/advisor/cohort/skill-gaps?course=LAW", headers=advisor).status_code == 400


class TestIntakeTemplate:
    def test_template_has_course_drop_down(self, env, advisor, admin, alice):
        from openpyxl import load_workbook
        c = env["client"]
        assert c.get("/api/intake/template", headers=alice).status_code == 403
        assert c.get("/api/intake/template", headers=admin).status_code == 200
        r = c.get("/api/intake/template", headers=advisor)
        assert r.status_code == 200 and r.headers["content-type"].startswith("application/vnd.openxmlformats")
        ws = load_workbook(io.BytesIO(r.content)).active
        assert [cell.value for cell in ws[1]] == ["Full Name", "Phone Number", "Personal Email", "Course", "Programme"]
        dv, prog_dv = ws.data_validations.dataValidation
        # Programme: the list of the course chosen in the same row (named ranges on a hidden sheet)
        assert prog_dv.formula1 == 'INDIRECT(SUBSTITUTE($D2," ","_"))' and "E2:E1001" in str(prog_dv.sqref)
        book = load_workbook(io.BytesIO(r.content))
        assert set(book.defined_names) == {"Computer_Science", "Information_and_Communications_Technology"}
        assert book["Lists"].sheet_state == "hidden"
        assert dv.type == "list" and dv.formula1 == '"Computer Science,Information and Communications Technology"'        # Dean of Computing
        assert "D2:D1001" in str(dv.sqref)
        assert [row[3].value for row in ws.iter_rows(min_row=2)] == ["Computer Science", "Information and Communications Technology"]
        every = load_workbook(io.BytesIO(c.get("/api/intake/template", headers=admin).content)).active
        assert every.data_validations.dataValidation[0].formula1 == '"Computer Science,Information and Communications Technology,Business"'
        # The filled-in template uploads as it is
        assert upload(env, advisor, "template.xlsx", r.content,
                      "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet").json()["students"] == 2


# ── Advisors: Dean / Head of Programme, courses, staff IDs ───
def staff_account(env, email, name, position, courses):
    """An advisor with a position, signed in (password not temporary)"""
    pw = bcrypt.hashpw(PASSWORD.encode(), bcrypt.gensalt()).decode()
    env["conn"].execute("INSERT INTO users (email, password_hash, role, full_name, advisor_position, advisor_courses) "
                        "VALUES (%s, %s, 'advisor', %s, %s, %s) ON CONFLICT (email) DO NOTHING",
                        (email, pw, name, position, courses))
    return token(env, email)


class TestAdvisorScope:
    def test_staff_numbers_in_the_database(self, env, biz_dean):
        conn = env["conn"]
        row = lambda e: conn.execute("SELECT advisor_position, advisor_courses, staff_number FROM users "
                                     "WHERE email = %s", (e,)).fetchone()
        pos, courses, number = row("adv@advisor.newinti.edu.my")     # inserted without a position
        assert (pos, courses) == ("dean", ["CS", "ICT"]) and number.startswith("10") and len(number) == 5
        assert row("dean.biz@advisor.newinti.edu.my")[2].startswith("20")
        assert conn.execute("SELECT count(*) FROM users WHERE role <> 'advisor' AND (staff_number IS NOT NULL "
                            "OR advisor_position IS NOT NULL)").fetchone()[0] == 0
        with pytest.raises(psycopg.errors.CheckViolation):           # a Dean of Computing ID must start with 10
            conn.execute("UPDATE users SET staff_number = '11999' WHERE email = 'adv@advisor.newinti.edu.my'")
        with pytest.raises(psycopg.errors.CheckViolation):           # no such position
            conn.execute("UPDATE users SET advisor_position = 'hop', advisor_courses = '{CS,ICT}' "
                         "WHERE email = 'adv@advisor.newinti.edu.my'")
        # One Dean per office: a second one is refused; without a position the next advisor is HOP CS
        with pytest.raises(psycopg.errors.UniqueViolation):
            conn.execute("INSERT INTO users (email, password_hash, role, full_name, advisor_position, advisor_courses) "
                         "VALUES ('dean2@advisor.newinti.edu.my', 'x', 'advisor', 'D2', 'dean', '{CS,ICT}')")
        extra = conn.execute("INSERT INTO users (email, password_hash, role, full_name) VALUES "
                             "('plain@advisor.newinti.edu.my', 'x', 'advisor', 'Plain') "
                             "RETURNING advisor_position, advisor_courses, staff_number").fetchone()
        assert extra[0] == "hop" and extra[1] == ["CS"] and extra[2].startswith("11")

    def test_admin_creates_and_edits_advisors(self, env, admin, biz_dean):
        c = env["client"]
        body = lambda email, **kw: {"email": email, "full_name": "Dr Tan", "role": "advisor",
                                    "personal_email": "p." + email.split("@")[0] + "@example.com", **kw}
        create = lambda b: c.post("/api/admin/users", headers=admin, json=b)

        assert create(body("a1@advisor.newinti.edu.my")).status_code == 400                     # no position
        assert create(body("a1@advisor.newinti.edu.my", advisor_position="dean",
                           advisor_courses=["CS"])).status_code == 400                          # Dean of CS only: no
        assert create(body("a1@advisor.newinti.edu.my", advisor_position="hop",
                           advisor_courses=["CS", "ICT"])).status_code == 400                   # HOP of two: no
        assert create({**body("x@admin.newinti.edu.my", advisor_position="hop", advisor_courses=["CS"]),
                       "role": "admin"}).status_code == 400                                     # admins have none

        suggested = c.get("/api/admin/staff-number?position=hop&courses=ICT", headers=admin).json()
        assert suggested["staff_number"].startswith("12") and suggested["position_label"] == "Head of Programme, Information and Communications Technology"
        hop = create(body("hop.ict@advisor.newinti.edu.my", advisor_position="hop", advisor_courses=["ICT"]))
        assert hop.status_code == 201, hop.text
        hop = hop.json()
        assert hop["staff_number"] == suggested["staff_number"] and hop["advisor_courses"] == ["ICT"]
        assert create(body("a2@advisor.newinti.edu.my", advisor_position="dean", advisor_courses=["ICT", "CS"],
                           staff_number="11001")).status_code == 400                            # wrong prefix
        assert create(body("a2@advisor.newinti.edu.my", advisor_position="hop", advisor_courses=["ICT"],
                           staff_number=hop["staff_number"])).status_code == 400                # taken
        assert create(body("a2@advisor.newinti.edu.my", advisor_position="hop", advisor_courses=["ICT"],
                           staff_number="1234")).status_code == 400                             # not 5 digits
        # One Dean per office: Computing already has one (the default test advisor)
        second = create(body("dean.comp@advisor.newinti.edu.my", advisor_position="dean", advisor_courses=["ICT", "CS"]))
        assert second.status_code == 400 and "Advisor A is already the Dean, Computing" in second.json()["detail"]
        info = c.get("/api/admin/staff-number?position=dean&courses=CS&courses=ICT", headers=admin).json()
        assert info["current_dean"]["full_name"] == "Advisor A"
        # Business: the old Dean is deactivated first, then the new one can be added
        old = env["conn"].execute("SELECT user_id::text FROM users WHERE email = 'dean.biz@advisor.newinti.edu.my'").fetchone()[0]
        biz = lambda: create(body("dean.biz2@advisor.newinti.edu.my", advisor_position="dean",
                                  advisor_courses=["BUSINESS"], staff_number="20777"))
        assert biz().status_code == 400
        assert c.patch(f"/api/admin/users/{old}", headers=admin, json={"is_active": False}).status_code == 200
        dean = biz()
        assert dean.status_code == 201 and dean.json()["advisor_courses"] == ["BUSINESS"]
        assert dean.json()["position_label"] == "Dean, Business" and dean.json()["staff_number"] == "20777"
        assert c.patch(f"/api/admin/users/{old}", headers=admin, json={"is_active": True}).status_code == 400
        # Edit: rename, then move to another course (a new staff ID with the right prefix is given)
        patch = lambda b: c.patch(f"/api/admin/users/{hop['user_id']}", headers=admin, json=b)
        # A Head of Programme cannot be moved into a Dean post that is taken
        assert patch({"advisor_position": "dean", "advisor_courses": ["CS", "ICT"]}).status_code == 400
        assert patch({"full_name": "Dr Tan Wei"}).json()["full_name"] == "Dr Tan Wei"
        moved = patch({"advisor_courses": ["CS"]}).json()
        assert moved["advisor_courses"] == ["CS"] and moved["staff_number"].startswith("11")
        assert patch({"staff_number": "11555"}).json()["staff_number"] == "11555"
        assert patch({"staff_number": "21555"}).status_code == 400
        alice = env["ids"]["alice"][0]
        assert c.patch(f"/api/admin/users/{alice}", headers=admin, json={"staff_number": "11556"}).status_code == 400
        found = c.get("/api/admin/users?search=11555", headers=admin).json()["users"]
        assert [u["email"] for u in found] == ["hop.ict@advisor.newinti.edu.my"]
        assert c.post(f"/api/admin/users/{hop['user_id']}/reset-password", headers=admin,
                      json={"password": "Reset#2026"}).status_code == 200

    def test_advisors_see_and_add_only_their_courses(self, env):
        c, conn = env["client"], env["conn"]
        pw = bcrypt.hashpw(PASSWORD.encode(), bcrypt.gensalt()).decode()
        for email, name, course in (("cs.kid@student.newinti.edu.my", "Cs Kid", "CS"),
                                    ("ict.kid@student.newinti.edu.my", "Ict Kid", "ICT")):
            uid = conn.execute("INSERT INTO users (email, password_hash, role, full_name, course) "
                               "VALUES (%s, %s, 'student', %s, %s) RETURNING user_id::text",
                               (email, pw, name, course)).fetchone()[0]
            conn.execute("INSERT INTO student_profiles (user_id) VALUES (%s)", (uid,))
        hop_cs = staff_account(env, "hop.cs@advisor.newinti.edu.my", "HOP CS", "hop", ["CS"])

        me = c.get("/api/auth/me", headers=hop_cs).json()
        assert me["advisor_courses"] == ["CS"] and me["staff_number"].startswith("11")
        assert me["position_label"] == "Head of Programme, Computer Science"

        students = {s["full_name"]: s for s in c.get("/api/advisor/students", headers=hop_cs).json()}
        assert "Cs Kid" in students and "Ict Kid" not in students
        assert all(s["course"] == "CS" for s in students.values())
        ict_uid = conn.execute("SELECT user_id::text FROM users WHERE email = 'ict.kid@student.newinti.edu.my'").fetchone()[0]
        assert c.get(f"/api/advisor/students/{ict_uid}", headers=hop_cs).status_code == 404
        assert c.get(f"/api/advisor/students/{students['Cs Kid']['user_id']}", headers=hop_cs).status_code == 200
        assert c.get("/api/advisor/cohort/overview?course=ICT", headers=hop_cs).status_code == 403
        assert c.get("/api/advisor/cohort/overview", headers=hop_cs).json()["courses"] == ["CS"]

        # New-student lists: rows of other courses are skipped with the reason
        csv = (b"Full Name,Personal Email,Course,Programme\n"
               b"Cs New,cs.new@example.com,Computer Science,Bachelor of Computer Science (Software Engineering)\n"
               b"Ict New,ict.new@example.com,ICT,Bachelor of Information and Communications Technology (Web and Mobile Development)\n")
        r = upload(env, hop_cs, "mine.csv", csv)
        assert r.status_code == 201 and r.json()["students"] == 1
        assert r.json()["skipped"][0]["line"] == 3 and "Computer Science" in r.json()["skipped"][0]["error"]


class TestSkillsAndProgrammes:
    def test_skill_list_follows_the_course(self, env):
        c, conn = env["client"], env["conn"]
        biz = course_student(env, "gina@student.newinti.edu.my", "Gina", "BUSINESS")
        cs = course_student(env, "hari@student.newinti.edu.my", "Hari", "CS")
        names = lambda h: {s["skill_name"] for s in c.get("/api/students/skills-library", headers=h).json()}
        biz_skills, cs_skills = names(biz), names(cs)
        assert {"Financial Modelling", "Payroll", "Excel", "Power BI", "Communication"} <= biz_skills
        assert not {"Docker", "Kubernetes", "React", "Metasploit"} & biz_skills          # IT-only skills
        assert {"Docker", "React", "Excel", "Communication", "Negotiation"} <= cs_skills
        assert not {"Financial Modelling", "Payroll"} & cs_skills          # Business-only skills
        docker = conn.execute("SELECT skill_id::text FROM skills_library WHERE skill_name = 'Docker'").fetchone()[0]
        r = c.post("/api/students/skills", headers=biz, json={"skill_id": docker})
        assert r.status_code == 400 and "course" in r.json()["detail"]

    def test_admin_skill_courses(self, env, admin):
        c = env["client"]
        biz = course_student(env, "gina@student.newinti.edu.my", "Gina", "BUSINESS")
        cs = course_student(env, "hari@student.newinti.edu.my", "Hari", "CS")
        names = lambda h: {s["skill_name"] for s in c.get("/api/students/skills-library", headers=h).json()}
        new = c.post("/api/admin/skills", headers=admin, json={"skill_name": "Retail Operations", "category": "Technical",
                                                               "courses": ["BUSINESS"]})
        assert new.status_code == 201 and new.json()["courses"] == ["BUSINESS"]
        assert "Retail Operations" in names(biz) and "Retail Operations" not in names(cs)
        upd = c.patch(f"/api/admin/skills/{new.json()['skill_id']}", headers=admin, json={"courses": ["ICT", "CS"]})
        assert upd.status_code == 200 and upd.json()["courses"] == ["CS", "ICT"]
        assert "Retail Operations" in names(cs) and "Retail Operations" not in names(biz)
        assert c.patch(f"/api/admin/skills/{new.json()['skill_id']}", headers=admin, json={"courses": []}).status_code == 422
        listed = {s["skill_name"]: s for s in c.get("/api/admin/skills", headers=admin).json()}
        assert listed["Docker"]["courses"] == ["CS", "ICT"] and listed["Payroll"]["courses"] == ["BUSINESS"]
        # Without courses an admin skill is for every course
        every = c.post("/api/admin/skills", headers=admin, json={"skill_name": "Public Speaking Club", "category": "Soft"})
        assert every.json()["courses"] == ["BUSINESS", "CS", "ICT"]

    def test_programmes_and_admin_change(self, env, admin, advisor):
        c, conn = env["client"], env["conn"]
        biz = c.get("/api/students/programmes?course=business", headers=advisor).json()
        assert [p["programme_name"] for p in biz] == ["Bachelor of Business (Business Analytics)", "Bachelor of Business (Marketing)",
                                                      "Bachelor of Business (Finance)", "Bachelor of Business (Human Resource Management)"]
        assert {p["course"] for p in c.get("/api/students/programmes", headers=advisor).json()} == {"CS", "BUSINESS", "ICT"}
        uid = conn.execute("SELECT user_id::text FROM users WHERE email = 'gina@student.newinti.edu.my'").fetchone()[0]
        patch = lambda body, u=uid: c.patch(f"/api/admin/users/{u}", headers=admin, json=body)
        assert patch({"programme": "Bachelor of Computer Science (Software Engineering)"}).status_code == 400          # not a Business programme
        r = patch({"programme": "bachelor of business (finance)"})
        assert r.status_code == 200 and r.json()["programme"] == "Bachelor of Business (Finance)"
        adv = conn.execute("SELECT user_id::text FROM users WHERE email = 'adv@advisor.newinti.edu.my'").fetchone()[0]
        assert patch({"programme": "Bachelor of Business (Marketing)"}, adv).status_code == 400     # only students

    def test_intake_programme_edit_and_account(self, env, admin):
        c, conn = env["client"], env["conn"]
        csv = f"Full Name,Personal Email,Course,Programme\nIvy Lee,ivy.lee@example.com,ICT,Bachelor of Information and Communications Technology (Web and Mobile Development)\n".encode()
        bid = upload(env, admin, "ivy.csv", csv).json()["batch_id"]
        row = c.get(f"/api/intake/batches/{bid}", headers=admin).json()["rows"][0]
        edit = lambda body: c.patch(f"/api/intake/rows/{row['row_id']}", headers=admin, json=body)
        assert edit({"programme": "Bachelor of Business (Marketing)"}).status_code == 400            # not an ICT programme
        assert edit({"course": "BUSINESS"}).json()["programme"] is None    # old programme no longer fits
        failed = c.post(f"/api/intake/batches/{bid}/create-accounts", headers=admin, json={}).json()["failed"]
        assert "programme" in failed[0]["error"]
        assert edit({"programme": "Bachelor of Business (Marketing)"}).json()["programme"] == "Bachelor of Business (Marketing)"
        done = c.post(f"/api/intake/batches/{bid}/create-accounts", headers=admin, json={}).json()
        assert len(done["created"]) == 1
        number = done["created"][0]["student_number"]
        prog = conn.execute("SELECT sp.programme FROM student_profiles sp JOIN users u USING (user_id) "
                            "WHERE u.student_number = %s", (number,)).fetchone()[0]
        assert prog == "Bachelor of Business (Marketing)"
        mail = next(e for e in c.get("/api/intake/outbox", headers=admin).json()["emails"]
                    if e["to_address"] == "ivy.lee@example.com")
        assert "Programme:           Bachelor of Business (Marketing)" in mail["body"]


# ── CV upload: skills found in the CV ────────────────────────
CV_TEXT = ("Jamie Tan - Computer Science student. Skills: Python, SQL and Docker. "
           "Built a React dashboard and did payroll for a student club. Strong communication skills.")


def cv_docx(text):
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("[Content_Types].xml", "<Types/>")
        z.writestr("word/document.xml",
                   '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
                   + "".join(f"<w:p><w:r><w:t>{line}</w:t></w:r></w:p>" for line in text.split(". "))
                   + "</w:body></w:document>")
    return buf.getvalue()


def cv_pdf(text):
    """A one-page PDF with a real text layer (hand-built, no extra package)"""
    stream = f"BT /F1 9 Tf 20 750 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for i, o in enumerate(objs, 1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + o + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1))
    for off in offsets:
        out.write(b"%010d 00000 n \n" % off)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref))
    return out.getvalue()


class TestCv:
    DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

    def upload(self, env, h, name, data, mime):
        return env["client"].post("/api/students/cv?method=keyword", headers=h, files={"file": (name, data, mime)})

    def test_upload_finds_library_skills_of_the_course(self, env):
        c, conn = env["client"], env["conn"]
        h = course_student(env, "jamie@student.newinti.edu.my", "Jamie", "CS")
        assert c.get("/api/students/cv", headers=h).status_code == 404
        r = self.upload(env, h, "jamie_cv.docx", cv_docx(CV_TEXT), self.DOCX)
        assert r.status_code == 201, r.text
        body = r.json()
        found = {s["skill_name"]: s for s in body["skills"]}
        assert {"Python", "SQL", "Docker", "React", "Communication"} <= set(found)
        assert body["method"] == "keyword" and body["file_name"] == "jamie_cv.docx"
        # Payroll is a Business skill: found, but CS students cannot add it
        assert found["Payroll"]["in_course"] is False and found["Python"]["in_course"] is True
        assert "Python" in found["Python"]["context"]
        # Nothing is added until the student confirms
        assert c.get("/api/students/skills", headers=h).json() == []

        picks = [{"skill_id": found[n]["skill_id"], "proficiency_level": "advanced" if n == "Python" else "intermediate"}
                 for n in ("Python", "SQL", "Payroll")]
        added = c.post("/api/students/cv/skills", headers=h, json={"skills": picks}).json()
        assert sorted(added["added"]) == ["Python", "SQL"]
        assert added["skipped"][0]["skill_name"] == "Payroll" and "course" in added["skipped"][0]["reason"]
        mine = {s["skill_name"]: s for s in c.get("/api/students/skills", headers=h).json()}
        assert mine["Python"]["proficiency_level"] == "advanced"
        again = c.post("/api/students/cv/skills", headers=h, json={"skills": picks[:1]}).json()
        assert again["added"] == [] and again["skipped"][0]["reason"] == "already added"
        # The stored CV can be read again; added skills are marked
        stored = c.get("/api/students/cv?method=keyword", headers=h).json()
        assert {s["skill_name"]: s["already_added"] for s in stored["skills"]}["Python"] is True
        dl = c.get("/api/students/cv/file", headers=h)
        assert dl.status_code == 200 and dl.content == cv_docx(CV_TEXT)
        # Delete: the CV goes, the skills stay
        assert c.delete("/api/students/cv", headers=h).status_code == 200
        assert c.get("/api/students/cv", headers=h).status_code == 404
        assert {"Python", "SQL"} <= {s["skill_name"] for s in c.get("/api/students/skills", headers=h).json()}

    def test_pdf_cv_and_bad_files(self, env, advisor):
        c = env["client"]
        h = course_student(env, "kim@student.newinti.edu.my", "Kim", "BUSINESS")
        r = self.upload(env, h, "kim.pdf", cv_pdf("Skills: Excel, Payroll, Recruitment and Docker."), "application/pdf")
        assert r.status_code == 201, r.text
        found = {s["skill_name"]: s["in_course"] for s in r.json()["skills"]}
        assert found["Payroll"] is True and found["Recruitment"] is True and found["Docker"] is False
        # A new upload replaces the old CV
        r2 = self.upload(env, h, "kim2.docx", cv_docx("Experienced in Excel and negotiation with suppliers"), self.DOCX)
        assert r2.status_code == 201 and r2.json()["file_name"] == "kim2.docx"
        assert env["conn"].execute("SELECT count(*) FROM student_cvs sc JOIN student_profiles sp USING (profile_id) "
                                   "JOIN users u USING (user_id) WHERE u.email = 'kim@student.newinti.edu.my'"
                                   ).fetchone()[0] == 1
        assert self.upload(env, h, "cv.png", b"\x89PNG\r\n\x1a\n" + b"0" * 50, "image/png").status_code == 400
        assert self.upload(env, h, "cv.pdf", b"not really a pdf", "application/pdf").status_code == 400
        assert self.upload(env, h, "empty.pdf", cv_pdf(""), "application/pdf").status_code == 400   # no text
        big = cv_docx("x") + b"0" * (5 * 1024 * 1024)
        assert self.upload(env, h, "big.docx", big, self.DOCX).status_code == 413
        assert c.post("/api/students/cv?method=nope", headers=h,
                      files={"file": ("a.docx", cv_docx(CV_TEXT), self.DOCX)}).status_code == 400
        # Only students have a CV
        assert c.get("/api/students/cv", headers=advisor).status_code == 403


# ── Advisor dashboard: charts, progress, notes, CV ───────────
class TestAdvisorFeatures:
    @staticmethod
    def biz(env):
        """A Business advisor (the biz_dean fixture is deactivated by an earlier test)"""
        return staff_account(env, "hop.biz.features@advisor.newinti.edu.my", "HOP Business", "hop", ["BUSINESS"])

    def test_overview_charts_and_filters(self, env, advisor, alice):
        c = env["client"]
        run(env, alice, role_id=env["da"])
        o = c.get("/api/advisor/cohort/overview", headers=advisor).json()
        assert sum(b["students"] for b in o["distribution"]) == o["students_assessed"]
        assert [b["band"] for b in o["distribution"]][-1] == "90-100" and len(o["distribution"]) == 10
        sc = o["status_counts"]
        assert sc["on_track"] + sc["developing"] + sc["needs_support"] + sc["not_assessed"] == o["total_students"]
        assert sc["needs_support"] == o["students_at_risk"]
        assert sum(r["students"] for r in o["target_roles"]) == o["total_students"]
        assert {"completed", "in_progress", "students_completed"} <= set(o["activities"])
        # Skill gaps for one assessed role only
        da = c.get(f"/api/advisor/cohort/skill-gaps?role_id={env['da']}", headers=advisor)
        assert da.status_code == 200 and len(da.json()) > 0
        assert c.get("/api/advisor/cohort/skill-gaps?role_id=nope", headers=advisor).status_code == 400
        # Student list columns used for sorting and filtering
        row = next(s for s in c.get("/api/advisor/students", headers=advisor).json() if s["full_name"] == "Alice")
        assert {"activities_completed", "note_count", "cgpa", "year_of_study", "profile_complete"} <= set(row)

    def test_cohort_progress(self, env, advisor):
        biz_dean = self.biz(env)
        c = env["client"]
        r = c.get("/api/advisor/cohort/progress?course=CS", headers=advisor)
        assert r.status_code == 200
        body = r.json()
        assert {"completed", "in_progress", "saved", "students_active", "students_completed"} <= set(body["totals"])
        assert isinstance(body["top_resources"], list) and isinstance(body["students"], list)
        assert all(s["course"] == "CS" for s in body["students"])
        assert c.get("/api/advisor/cohort/progress?course=CS", headers=biz_dean).status_code == 403

    def test_notes_and_ai_review(self, env, advisor, alice):
        biz_dean = self.biz(env)
        c, conn = env["client"], env["conn"]
        uid = env["ids"]["alice"][0]
        run(env, alice, role_id=env["da"])
        aid = conn.execute("SELECT assessment_id::text FROM assessments WHERE profile_id = %s ORDER BY created_at DESC LIMIT 1",
                           (env["ids"]["alice"][1],)).fetchone()[0]
        notes = f"/api/advisor/students/{uid}/notes"
        assert c.post(notes, headers=advisor, json={"body": ""}).status_code == 422
        n1 = c.post(notes, headers=advisor, json={"body": "Talked about an internship in data analytics."})
        assert n1.status_code == 201 and n1.json()["note_type"] == "general" and n1.json()["can_edit"] is True
        # A review of the AI result needs the assessment and a verdict; the assessment must be the student's
        assert c.post(notes, headers=advisor, json={"body": "x", "note_type": "ai_review"}).status_code == 400
        other = conn.execute("SELECT assessment_id::text FROM assessments WHERE profile_id <> %s LIMIT 1",
                             (env["ids"]["alice"][1],)).fetchone()
        if other:
            assert c.post(notes, headers=advisor, json={"body": "x", "note_type": "ai_review", "ai_verdict": "agree",
                                                        "assessment_id": other[0]}).status_code == 400
        n2 = c.post(notes, headers=advisor, json={"body": "SQL is stronger than the score shows (portfolio project).",
                                                  "note_type": "ai_review", "ai_verdict": "disagree", "assessment_id": aid})
        assert n2.status_code == 201 and n2.json()["ai_verdict"] == "disagree" and n2.json()["assessment_role"] == "Data Analyst"
        assert c.post(notes, headers=advisor, json={"body": "x", "ai_verdict": "agree"}).status_code == 400
        listed = c.get(notes, headers=advisor).json()
        assert [n["note_id"] for n in listed][:2] == [n2.json()["note_id"], n1.json()["note_id"]]   # newest first
        assert c.get(f"/api/advisor/students/{uid}", headers=advisor).json()["notes"][0]["note_id"] == n2.json()["note_id"]
        assert next(s for s in c.get("/api/advisor/students", headers=advisor).json() if s["user_id"] == uid)["note_count"] >= 2

        # Another advisor of the course sees the notes but cannot change them; other courses cannot see them
        hop = staff_account(env, "hop.cs.notes@advisor.newinti.edu.my", "HOP CS Notes", "hop", ["CS"])
        seen = c.get(notes, headers=hop).json()
        assert {n["note_id"] for n in seen} >= {n1.json()["note_id"]} and all(not n["can_edit"] for n in seen)
        assert c.patch(f"/api/advisor/notes/{n1.json()['note_id']}", headers=hop, json={"body": "hacked"}).status_code == 403
        assert c.delete(f"/api/advisor/notes/{n1.json()['note_id']}", headers=hop).status_code == 403
        assert c.get(notes, headers=biz_dean).status_code == 404
        assert c.get(notes, headers=alice).status_code == 403                       # never shown to the student

        # The author edits and deletes
        upd = c.patch(f"/api/advisor/notes/{n2.json()['note_id']}", headers=advisor, json={"ai_verdict": "agree"})
        assert upd.status_code == 200 and upd.json()["ai_verdict"] == "agree"
        assert c.patch(f"/api/advisor/notes/{n1.json()['note_id']}", headers=advisor,
                       json={"ai_verdict": "agree"}).status_code == 400              # a general note has no verdict
        assert c.delete(f"/api/advisor/notes/{n1.json()['note_id']}", headers=advisor).status_code == 200
        assert c.delete(f"/api/advisor/notes/{n1.json()['note_id']}", headers=advisor).status_code == 404
        actions = {r[0] for r in conn.execute("SELECT DISTINCT action FROM audit_log").fetchall()}
        assert {"advisor.note_create", "advisor.note_update", "advisor.note_delete"} <= actions

    def test_advisor_downloads_student_cv(self, env, advisor):
        biz_dean = self.biz(env)
        c, conn = env["client"], env["conn"]
        h = course_student(env, "lena@student.newinti.edu.my", "Lena", "CS")
        uid = conn.execute("SELECT user_id::text FROM users WHERE email = 'lena@student.newinti.edu.my'").fetchone()[0]
        c.get("/api/students/skills", headers=h)
        conn.execute("INSERT INTO student_profiles (user_id) VALUES (%s) ON CONFLICT DO NOTHING", (uid,))
        assert c.get(f"/api/advisor/students/{uid}/cv", headers=advisor).status_code == 404     # none yet
        data = cv_docx(CV_TEXT)
        assert c.post("/api/students/cv?method=keyword", headers=h,
                      files={"file": ("lena.docx", data, TestCv.DOCX)}).status_code == 201
        assert c.get(f"/api/advisor/students/{uid}", headers=advisor).json()["cv"]["file_name"] == "lena.docx"
        dl = c.get(f"/api/advisor/students/{uid}/cv", headers=advisor)
        assert dl.status_code == 200 and dl.content == data
        assert c.get(f"/api/advisor/students/{uid}/cv", headers=biz_dean).status_code == 404
        assert conn.execute("SELECT count(*) FROM audit_log WHERE action = 'advisor.view_cv' AND target_id = %s",
                            (uid,)).fetchone()[0] == 1


# ── Work experience ──────────────────────────────────────────
class TestExperience:
    def test_add_edit_delete_and_scores(self, env, advisor):
        c, conn = env["client"], env["conn"]
        h = course_student(env, "maya@student.newinti.edu.my", "Maya", "CS")
        uid = conn.execute("SELECT user_id::text FROM users WHERE email = 'maya@student.newinti.edu.my'").fetchone()[0]
        ex = "/api/students/experiences"
        assert c.get(ex, headers=h).json() == []
        job = {"job_title": "Data Analyst Intern", "company": "Acme Sdn Bhd", "experience_type": "internship",
               "location": "Kuala Lumpur", "start_date": "2026-01-05", "end_date": "2026-04-03",
               "description": "Built Power BI dashboards."}
        assert c.post(ex, headers=h, json={**job, "start_date": "2099-01-01", "end_date": None}).status_code == 422
        assert c.post(ex, headers=h, json={**job, "end_date": "2025-12-01"}).status_code == 422
        assert c.post(ex, headers=h, json={**job, "experience_type": "astronaut"}).status_code == 422
        r = c.post(ex, headers=h, json=job)
        assert r.status_code == 201
        eid = r.json()["experience_id"]
        c.post(ex, headers=h, json={"job_title": "Tutor", "company": "Campus Club", "experience_type": "volunteer",
                                    "start_date": "2025-06-01"})                       # still going on
        listed = c.get(ex, headers=h).json()
        assert [e["job_title"] for e in listed] == ["Tutor", "Data Analyst Intern"]    # current first
        assert c.put(f"{ex}/{eid}", headers=h, json={**job, "job_title": "BI Intern"}).status_code == 200
        assert "BI Intern" in {e["job_title"] for e in c.get(ex, headers=h).json()}
        assert conn.execute("SELECT internship_count FROM student_profiles sp JOIN users u USING (user_id) "
                            "WHERE u.user_id = %s", (uid,)).fetchone()[0] == 1
        # Another student cannot change it
        other = course_student(env, "nora@student.newinti.edu.my", "Nora", "CS")
        assert c.put(f"{ex}/{eid}", headers=other, json=job).status_code == 404
        assert c.delete(f"{ex}/{eid}", headers=other).status_code == 404
        # Advisors see it on the student's page
        detail = c.get(f"/api/advisor/students/{uid}", headers=advisor).json()
        assert {e["job_title"] for e in detail["experiences"]} == {"BI Intern", "Tutor"}
        # Paid experience counts as industry experience; volunteering does not
        sid = conn.execute("SELECT skill_id::text FROM skills_library WHERE skill_name = 'SQL'").fetchone()[0]
        c.post("/api/students/skills", headers=h, json={"skill_id": sid})
        result = run(env, h, role_id=env["da"]).json()
        assert float(result["dimension_scores"]["industry_experience"]) == 50.0
        assert c.delete(f"{ex}/{eid}", headers=h).status_code == 200
        assert [e["job_title"] for e in c.get(ex, headers=h).json()] == ["Tutor"]


# ── Career paths (FR-15) ─────────────────────────────────────
class TestCareerPaths:
    def test_paths_for_the_course(self, env, advisor):
        c, conn = env["client"], env["conn"]
        h = course_student(env, "omar@student.newinti.edu.my", "Omar", "CS")
        empty = c.get("/api/students/career-paths", headers=h).json()
        assert empty["skill_count"] == 0 and all(p["match_pct"] == 0 for p in empty["paths"])
        for skill in ("SQL", "Python", "Excel", "Power BI", "Tableau", "Data Visualisation", "Communication"):
            sid = conn.execute("SELECT skill_id::text FROM skills_library WHERE skill_name = %s", (skill,)).fetchone()[0]
            c.post("/api/students/skills", headers=h, json={"skill_id": sid})
        c.put("/api/students/profile", headers=h, json={"target_role": "Data Analyst"})
        body = c.get("/api/students/career-paths", headers=h).json()
        names = [p["role_name"] for p in body["paths"]]
        assert {"Software Developer", "Data Scientist", "Data Analyst"} <= set(names)
        assert "Financial Analyst" not in names                                     # Business role
        assert [p["match_pct"] for p in body["paths"]] == sorted((p["match_pct"] for p in body["paths"]), reverse=True)
        da = next(p for p in body["paths"] if p["role_name"] == "Data Analyst")
        assert da["is_target_role"] is True and da["match_pct"] > 0
        assert "SQL" in da["matched_skills"] and da["required"]["have"] >= 2
        assert all(m["skill_name"] not in da["matched_skills"] for m in da["missing_skills"])
        assert all(s["importance"] == "required" for s in da["next_steps"]) and len(da["next_steps"]) <= 3
        assert any(s["resource"] and s["resource"]["url"] for s in da["next_steps"])
        biz = course_student(env, "pia@student.newinti.edu.my", "Pia", "BUSINESS")
        assert {p["role_name"] for p in c.get("/api/students/career-paths", headers=biz).json()["paths"]} == \
            {"Business Analyst", "Marketing Executive", "Financial Analyst", "HR Executive"}
        assert c.get("/api/students/career-paths", headers=advisor).status_code == 403


# ── System settings (FR-16) ──────────────────────────────────
class TestSettings:
    def test_admin_changes_settings(self, env, admin, advisor, alice):
        c, conn = env["client"], env["conn"]
        defaults = c.get("/api/settings", headers=alice).json()
        assert defaults == {"semester_starts": ["01-01", "05-01", "09-01"], "readiness_on_track": 70,
                            "readiness_needs_support": 40, "max_upload_mb": 5, "max_evidence_per_item": 5}
        assert c.get("/api/admin/settings", headers=advisor).status_code == 403
        assert c.put("/api/admin/settings", headers=advisor, json={"values": {"max_upload_mb": 1}}).status_code == 403
        view = c.get("/api/admin/settings", headers=admin).json()
        assert {x["key"] for x in view["settings"]} == set(defaults) and view["email"]["configured"] is False
        put = lambda values: c.put("/api/admin/settings", headers=admin, json={"values": values})
        assert put({"readiness_needs_support": 80}).status_code == 400              # not below On Track
        assert put({"max_upload_mb": 6}).status_code == 400                         # database limit is 5 MB
        assert put({"semester_starts": ["13-01"]}).status_code == 400
        assert put({"semester_starts": []}).status_code == 400
        assert put({"colour": "blue"}).status_code == 400
        assert put({"readiness_on_track": 2.5}).status_code == 400
        try:
            r = put({"readiness_on_track": 60, "readiness_needs_support": 30, "max_evidence_per_item": 1, "max_upload_mb": 1})
            assert r.status_code == 200
            row = next(x for x in r.json()["settings"] if x["key"] == "max_upload_mb")
            assert row["value"] == 1 and row["updated_by"] == "Admin A"
            assert c.get("/api/advisor/cohort/overview", headers=advisor).json()["bands"] == {"on_track": 60, "needs_support": 30}
            # Upload limits apply at once
            h = course_student(env, "quinn@student.newinti.edu.my", "Quinn", "CS")
            big = cv_docx("x") + b"0" * (1024 * 1024 + 10)
            r = c.post("/api/students/cv?method=keyword", headers=h, files={"file": ("big.docx", big, TestCv.DOCX)})
            assert r.status_code == 413 and "1 MB" in r.json()["detail"]
            pid = c.post("/api/students/projects", headers=h, json={"project_title": "Site"}).json()["project_id"]
            link = lambda: c.post(f"/api/students/projects/{pid}/links", headers=h, json={"url": "https://example.com/a"})
            assert link().status_code == 201
            r2 = link()
            assert r2.status_code == 400 and "at most 1" in r2.json()["detail"]
            assert conn.execute("SELECT count(*) FROM audit_log WHERE action = 'admin.settings_update'").fetchone()[0] >= 1
        finally:
            assert put({"readiness_on_track": 70, "readiness_needs_support": 40, "max_evidence_per_item": 5,
                        "max_upload_mb": 5}).status_code == 200
        assert c.get("/api/settings", headers=alice).json() == defaults

    def test_semester_starts_setting(self, env, admin):
        from datetime import date
        from routers.students import current_semester_start
        assert current_semester_start(date(2026, 8, 1), ((2, 15), (8, 1))) == date(2026, 8, 1)
        assert current_semester_start(date(2026, 1, 10), ((2, 15), (8, 1))) == date(2025, 8, 1)


# ── Email: test email, resend, fake mail server ──────────────
class FakeSMTP:
    sent, fail = [], False

    def __init__(self, host, port, timeout=None):
        self.host, self.port = host, port

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def starttls(self):
        pass

    def login(self, user, password):
        pass

    def send_message(self, msg):
        if FakeSMTP.fail:
            raise OSError("connection refused")
        FakeSMTP.sent.append(msg)


class TestEmail:
    def test_test_email_and_resend(self, env, admin, monkeypatch):
        import notifications
        c, conn = env["client"], env["conn"]
        assert c.post("/api/admin/settings/test-email", headers=admin, json={"to": "ops@example.com"}).status_code == 400
        # An account email stored while no mail server was set up
        pw = bcrypt.hashpw(PASSWORD.encode(), bcrypt.gensalt()).decode()
        uid = conn.execute("INSERT INTO users (email, password_hash, role, full_name, course, must_change_password) "
                           "VALUES ('rae@student.newinti.edu.my', %s, 'student', 'Rae', 'CS', TRUE) RETURNING user_id::text",
                           (pw,)).fetchone()[0]
        eid = conn.execute("INSERT INTO email_outbox (to_address, subject, body, status, related_user_id) VALUES "
                           "('rae.personal@example.com', 'Your SEAGAS student account', 'Temporary password:  Abc#1234x', "
                           "'not_sent', %s) RETURNING email_id::text", (uid,)).fetchone()[0]
        assert c.post(f"/api/intake/outbox/{eid}/resend", headers=admin).status_code == 400   # no mail server

        monkeypatch.setenv("SMTP_HOST", "smtp.example.com")
        monkeypatch.setenv("SMTP_FROM", "seagas@example.com")
        monkeypatch.setattr(notifications.smtplib, "SMTP", FakeSMTP)
        FakeSMTP.sent.clear()
        status = c.get("/api/admin/settings", headers=admin).json()["email"]
        assert status["configured"] and status["host"] == "smtp.example.com" and status["security"] == "starttls"
        assert "password" not in str(status).lower()
        r = c.post("/api/admin/settings/test-email", headers=admin, json={"to": "ops@example.com"})
        assert r.json() == {"status": "sent", "error": None} and FakeSMTP.sent[-1]["To"] == "ops@example.com"

        r = c.post(f"/api/intake/outbox/{eid}/resend", headers=admin)
        assert r.json()["status"] == "sent" and FakeSMTP.sent[-1]["To"] == "rae.personal@example.com"
        row = next(e for e in c.get("/api/intake/outbox", headers=admin).json()["emails"] if e["email_id"] == eid)
        assert row["status"] == "sent" and row["attempts"] == 2 and row["sent_at"] and row["password_still_temporary"] is True
        FakeSMTP.fail = True
        try:
            r = c.post(f"/api/intake/outbox/{eid}/resend", headers=admin)
            assert r.json()["status"] == "failed" and "connection refused" in r.json()["error"]
        finally:
            FakeSMTP.fail = False
        # After the person chose their own password the old temporary one must not be sent again
        conn.execute("UPDATE users SET must_change_password = FALSE WHERE user_id = %s", (uid,))
        r = c.post(f"/api/intake/outbox/{eid}/resend", headers=admin)
        assert r.status_code == 400 and "Reset their password" in r.json()["detail"]
        assert c.post(f"/api/intake/outbox/{eid}/resend", headers=course_student(env, "sam@student.newinti.edu.my", "Sam", "CS")).status_code == 403
        assert conn.execute("SELECT count(*) FROM audit_log WHERE action = 'admin.email_resend'").fetchone()[0] >= 2


class TestNewStudent:
    def test_new_student_without_profile_can_add_skills(self, env):
        c, conn = env["client"], env["conn"]
        pw = bcrypt.hashpw(PASSWORD.encode(), bcrypt.gensalt()).decode()
        conn.execute("INSERT INTO users (email, password_hash, role, full_name, course) "
                     "VALUES ('fresh@student.newinti.edu.my', %s, 'student', 'Fresh', 'ICT')", (pw,))
        h = token(env, "fresh@student.newinti.edu.my")
        # Everything the Profile page loads works without a profile
        assert len(c.get("/api/students/skills-library", headers=h).json()) > 100
        assert c.get("/api/students/skills", headers=h).json() == []
        assert c.get("/api/students/projects", headers=h).json() == []
        assert c.get("/api/students/certifications", headers=h).json() == []
        assert c.get("/api/students/profile", headers=h).status_code == 404
        # Adding a skill creates the (empty) profile
        sid = conn.execute("SELECT skill_id::text FROM skills_library WHERE skill_name = 'Python'").fetchone()[0]
        assert c.post("/api/students/skills", headers=h, json={"skill_id": sid}).status_code == 201
        assert c.get("/api/students/profile", headers=h).status_code == 200
        # Saving the Info tab afterwards still works, even though the page may send "create"
        info = {"programme": "Made Up", "university": "INTI", "year_of_study": 1, "target_role": "Data Analyst"}
        assert c.post("/api/students/profile", headers=h, json=info).status_code in (200, 201)
        p = c.get("/api/students/profile", headers=h).json()
        assert p["university"] == "INTI" and p["programme"] is None      # set by the university, not the student


# ── Batch assessments (data/run_assessments.py uses this) ────
def test_assessment_service_matches_endpoint(env, alice):
    import assessment_service
    from database import SessionLocal
    db = SessionLocal()
    try:
        profile = db.execute(assessment_service.text("SELECT * FROM student_profiles WHERE profile_id = :p"),
                             {"p": env["ids"]["alice"][1]}).fetchone()
        direct = assessment_service.run_for_profile(db, profile, role_id=env["da"], method="keyword")
    finally:
        db.close()
    via_api = run(env, alice, role_id=env["da"]).json()
    assert direct["readiness_score"] == via_api["readiness_score"]
    assert [m["match_status"] for m in direct["match_results"]] == [m["match_status"] for m in via_api["match_results"]]
