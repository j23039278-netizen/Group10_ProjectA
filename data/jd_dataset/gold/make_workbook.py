"""
Build the Excel workbook used to hand-label the gold JDs.

One sheet per JD: the JD text at the top (one row per line, wrapped) and the
label table below it, with columns skill_name, importance, evidence_text,
keep (Y/N) and note.

  Annotator A (default): all JDs in gold_sample.csv. The table is pre-filled
      from gold_labels_draft.csv (keyword baseline) with keep left empty, so
      every draft row must be marked Y or N by hand.
  Annotator B (--annotator B): the JDs given with --refs, or 2 per role drawn
      with seed 42. The table starts empty: B must not see the draft or A's
      labels.

skill_name has a drop-down of the canonical names in skills_library.csv
(free text such as "NEW: <term>" is still accepted after a warning);
importance has a required / preferred / bonus drop-down.

An existing workbook is never overwritten without --force, because it may
already hold a person's labels.

Usage (from repo root):
    python data/jd_dataset/gold/make_workbook.py                 # -> gold_workbook.xlsx (A)
    python data/jd_dataset/gold/make_workbook.py --annotator B   # -> gold_workbook_B.xlsx
Then read the result back with collect_workbook.py.
"""

import argparse
import math
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from gold_common import (DRAFT_PATH, IMPORTANCE, SAMPLE_PATH, SKILLS_PATH, TABLE_COLUMNS,
                         WORKBOOK_PATH, Catalogue, load_skills, pick_b_refs, read_csv)

LIST_SHEET = "skill_list"
README_SHEET = "README"
EXTRA_ROWS = 30  # empty table rows for skills added by hand
COL_WIDTHS = {"A": 32, "B": 13, "C": 85, "D": 7, "E": 30}
TEXT_CHARS_PER_LINE = 150  # rough wrap width of the merged A:E text rows
EVIDENCE_CHARS_PER_LINE = 80
LINE_HEIGHT = 15

BOLD = Font(bold=True)
TITLE_FONT = Font(bold=True, size=13)
HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
TEXT_FILL = PatternFill("solid", fgColor="F7F7F7")
TODO_FILL = PatternFill("solid", fgColor="FFF2CC")  # draft row not reviewed yet
DROP_FONT = Font(color="999999", strike=True)       # keep = N
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")

INSTRUCTIONS = {
    "A": ("Read the JD, then go down to the table. Mark every pre-filled row keep = Y "
          "(correct, fix the cells if needed) or N (wrong). Add missing skills on the empty "
          "rows with keep = Y. Yellow rows still need a Y or N."),
    "B": ("Read the JD, then go down to the table and add one row per skill "
          "(skill_name, importance, evidence_text). keep can stay empty; N removes a row."),
}


def set_text(cell, value):
    """Write a value as plain text (a leading '=' must not become a formula)."""
    cell.value = value
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


def lines_needed(text, width):
    return max(1, math.ceil(len(text) / width))


def add_list_sheet(wb, names):
    ws = wb.create_sheet(LIST_SHEET)
    for i, name in enumerate(names, start=1):
        set_text(ws.cell(row=i, column=1), name)
    ws.sheet_state = "hidden"
    wb.defined_names["skill_names"] = DefinedName(
        "skill_names", attr_text=f"{LIST_SHEET}!$A$1:$A${len(names)}")


def add_validations(ws, first, last):
    skill_dv = DataValidation(
        type="list", formula1="skill_names", allow_blank=True,
        showErrorMessage=True, errorStyle="warning",
        errorTitle="Not a canonical skill name",
        error="Pick a name from the list. For a skill the library does not have, "
              "type 'NEW: <term>' and click Yes.",
        showInputMessage=True, promptTitle="skill_name",
        prompt="Canonical name from the list, or NEW: <term>")
    imp_dv = DataValidation(type="list", formula1='"' + ",".join(IMPORTANCE) + '"',
                            allow_blank=True, showErrorMessage=True,
                            errorTitle="importance", error="Use required, preferred or bonus.")
    keep_dv = DataValidation(type="list", formula1='"Y,N"', allow_blank=True,
                             showErrorMessage=True, errorTitle="keep", error="Use Y or N.")
    for dv, col in ((skill_dv, "A"), (imp_dv, "B"), (keep_dv, "D")):
        ws.add_data_validation(dv)
        dv.add(f"{col}{first}:{col}{last}")


def add_jd_sheet(wb, jd, rows, annotator):
    """Write one JD sheet; returns (header_row, last_table_row)."""
    ws = wb.create_sheet(jd["jd_ref"])
    for col, width in COL_WIDTHS.items():
        ws.column_dimensions[col].width = width

    ws.merge_cells("A1:E1")
    set_text(ws["A1"], f"{jd['jd_ref']}  |  {jd['role_code']}  |  {jd['title']}  ({jd['seniority']})")
    ws["A1"].font = TITLE_FONT
    ws.merge_cells("A2:E2")
    ws["A2"].value = INSTRUCTIONS[annotator]
    ws["A2"].alignment = WRAP_TOP
    ws.row_dimensions[2].height = 2 * LINE_HEIGHT + 2

    r = 4
    ws.cell(row=r, column=1, value="Job description").font = BOLD
    r += 1
    for line in jd["raw_text"].split("\n"):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
        cell = ws.cell(row=r, column=1)
        set_text(cell, line)
        cell.alignment = WRAP_TOP
        cell.fill = TEXT_FILL
        ws.row_dimensions[r].height = (lines_needed(line, TEXT_CHARS_PER_LINE) * LINE_HEIGHT + 2
                                       if line.strip() else 8)
        r += 1

    header = r + 1
    for c, name in enumerate(TABLE_COLUMNS, start=1):
        cell = ws.cell(row=header, column=c, value=name)
        cell.font, cell.fill, cell.border = BOLD, HEADER_FILL, BOX
    for i, row in enumerate(rows, start=header + 1):
        for c, key in enumerate(("skill_name", "importance", "evidence_text"), start=1):
            set_text(ws.cell(row=i, column=c), row[key])
        ws.row_dimensions[i].height = (
            lines_needed(row["evidence_text"], EVIDENCE_CHARS_PER_LINE) * LINE_HEIGHT + 2)
    last = header + len(rows) + EXTRA_ROWS
    for i in range(header + 1, last + 1):
        for c in range(1, 6):
            cell = ws.cell(row=i, column=c)
            cell.alignment, cell.border = WRAP_TOP, BOX

    body = f"A{header + 1}:E{last}"
    if annotator == "A":
        ws.conditional_formatting.add(body, FormulaRule(
            formula=[f'AND($A{header + 1}<>"",$D{header + 1}="")'], fill=TODO_FILL))
    ws.conditional_formatting.add(body, FormulaRule(
        formula=[f'UPPER($D{header + 1})="N"'], font=DROP_FONT))
    add_validations(ws, header + 1, last)

    # Jump link from the top of the sheet to the table.
    link = ws.cell(row=3, column=1, value="Go to the label table")
    link.hyperlink = f"#'{jd['jd_ref']}'!A{header}"
    link.font = Font(color="0563C1", underline="single")
    return header, last


def add_readme(ws, annotator, entries):
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 50
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 16
    ws["A1"] = f"SEAGAS gold annotation - annotator {annotator}"
    ws["A1"].font = TITLE_FONT
    notes = [
        INSTRUCTIONS[annotator],
        "Follow data/jd_dataset/gold/ANNOTATION_GUIDE.md (what counts as a skill, importance cues).",
        "skill_name: pick a canonical name from the drop-down, or type NEW: <term> (click Yes on the warning).",
        "evidence_text: copy the shortest phrase from the JD text. End it with ? if unsure about importance.",
        "One row per skill per JD. Do not rename sheets or move the table header.",
    ]
    if annotator == "B":
        notes.append("Label on your own: do not look at the draft, annotator A's workbook or gold_labels.csv.")
    for i, note in enumerate(notes, start=2):
        ws.cell(row=i, column=1, value=note)
    top = len(notes) + 3
    status_col = "still to review" if annotator == "A" else "rows"
    for c, name in enumerate(("jd_ref", "role", "title", "keep = Y", status_col), start=1):
        cell = ws.cell(row=top, column=c, value=name)
        cell.font, cell.fill = BOLD, HEADER_FILL
    for i, (jd, header, last) in enumerate(entries, start=top + 1):
        ref = jd["jd_ref"]
        skills = f"'{ref}'!$A${header + 1}:$A${last}"
        keeps = f"'{ref}'!$D${header + 1}:$D${last}"
        cell = ws.cell(row=i, column=1, value=ref)
        cell.hyperlink = f"#'{ref}'!A{header}"
        cell.font = Font(color="0563C1", underline="single")
        ws.cell(row=i, column=2, value=jd["role_code"])
        set_text(ws.cell(row=i, column=3), jd["title"])
        ws.cell(row=i, column=4, value=f'=COUNTIF({keeps},"Y")')
        ws.cell(row=i, column=5, value=(f'=COUNTIFS({skills},"<>",{keeps},"")' if annotator == "A"
                                        else f"=COUNTA({skills})"))
    ws.freeze_panes = ws.cell(row=top + 1, column=1)


def build_workbook(jds, draft_rows, skill_names, annotator, out_path):
    """Write the workbook for `jds` (sample rows) and return its path."""
    wb = Workbook()
    readme = wb.active
    readme.title = README_SHEET
    add_list_sheet(wb, skill_names)
    by_ref = {}
    if annotator == "A":
        for row in draft_rows:
            by_ref.setdefault(row["jd_ref"], []).append(row)
    entries = []
    for jd in jds:
        header, last = add_jd_sheet(wb, jd, by_ref.get(jd["jd_ref"], []), annotator)
        entries.append((jd, header, last))
    add_readme(readme, annotator, entries)
    wb.move_sheet(LIST_SHEET, offset=len(wb.sheetnames))  # keep the hidden list last
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def main(argv=None):
    ap = argparse.ArgumentParser(description="Build the gold annotation workbook")
    ap.add_argument("--annotator", choices=["A", "B"], default="A")
    ap.add_argument("--refs", nargs="+", metavar="JD_REF",
                    help="JDs to include (B default: 2 per role, seed 42; A default: all)")
    ap.add_argument("--sample", default=str(SAMPLE_PATH))
    ap.add_argument("--draft", default=str(DRAFT_PATH))
    ap.add_argument("--skills", default=str(SKILLS_PATH))
    ap.add_argument("--out", help="output .xlsx (default: gold_workbook.xlsx / gold_workbook_B.xlsx)")
    ap.add_argument("--force", action="store_true", help="overwrite an existing workbook")
    args = ap.parse_args(argv)

    out = Path(args.out or WORKBOOK_PATH[args.annotator])
    if out.exists() and not args.force:
        sys.exit(f"{out} already exists and may hold labels. Use --force to overwrite it.")

    sample = read_csv(args.sample)
    known = {r["jd_ref"]: r for r in sample}
    refs = [ref for arg in args.refs for ref in arg.split(",") if ref] if args.refs else None
    if refs is None:
        refs = pick_b_refs(sample) if args.annotator == "B" else [r["jd_ref"] for r in sample]
    unknown = [ref for ref in refs if ref not in known]
    if unknown:
        sys.exit(f"Unknown jd_ref(s), not in {args.sample}: {', '.join(unknown)}")

    draft = read_csv(args.draft) if args.annotator == "A" else []
    names = Catalogue(load_skills(args.skills)).names
    jds = [known[ref] for ref in refs]
    build_workbook(jds, draft, names, args.annotator, out)
    n_draft = sum(1 for d in draft if d["jd_ref"] in set(refs))
    print(f"{out}: {len(jds)} JD sheets for annotator {args.annotator}"
          + (f", {n_draft} draft rows to review" if args.annotator == "A" else ", tables empty"))
    if args.annotator == "B":
        print("B refs: " + " ".join(refs))


if __name__ == "__main__":
    main()
