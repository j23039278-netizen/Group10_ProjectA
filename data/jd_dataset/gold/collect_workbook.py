"""
Read a filled gold workbook (from make_workbook.py) back into a label CSV.

  Annotator A: keeps the rows marked keep = Y. Rows marked N are dropped.
      A row with a skill but no Y/N is reported as "not reviewed" and dropped.
      Output: gold_labels.csv
  Annotator B: keeps every non-empty row (except keep = N).
      Output: gold_labels_B.csv

Checks on every kept row:
  errors   (the CSV is not written unless --skip-invalid drops these rows):
           unknown jd_ref, skill_name that is not a canonical library name or
           "NEW: <term>", NEW: term that is already in the library, bad
           importance, bad keep value
  warnings (reported only): evidence_text missing or not found in the JD text
           (case and whitespace ignored, a trailing '?' ignored)
Rows for the same skill in one JD are merged, keeping the strongest
importance (required > preferred > bonus) and its evidence.

The workbook is the source of truth: fix problems there and run this again.
Do not hand-edit the output CSV, it is overwritten on every run.

Usage (from repo root):
    python data/jd_dataset/gold/collect_workbook.py                 # A -> gold_labels.csv
    python data/jd_dataset/gold/collect_workbook.py --annotator B   # B -> gold_labels_B.csv
"""

import argparse
import sys
from dataclasses import dataclass, field

from openpyxl import load_workbook

from gold_common import (DRAFT_PATH, IMPORTANCE, IMPORTANCE_RANK, LABEL_FIELDS, LABELS_PATH,
                         NEW_PREFIX, SAMPLE_PATH, SKILLS_PATH, TABLE_COLUMNS, WORKBOOK_PATH,
                         Catalogue, evidence_in_text, load_skills, read_csv, skill_key, write_csv)

KEEP_YES = {"y", "yes"}
KEEP_NO = {"n", "no"}


def cell_str(value):
    return "" if value is None else str(value).strip()


def read_workbook(path):
    """{sheet_title: [row dict with excel_row + TABLE_COLUMNS]} for every sheet
    that has the label table header (README and the hidden list are skipped)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = {}
    for ws in wb.worksheets:
        rows, header = [], None
        for i, values in enumerate(ws.iter_rows(max_col=len(TABLE_COLUMNS), values_only=True), start=1):
            values = [cell_str(v) for v in values] + [""] * (len(TABLE_COLUMNS) - len(values))
            if header is None:
                if [v.lower() for v in values] == TABLE_COLUMNS:
                    header = i
                continue
            if any(values):
                rows.append({"excel_row": i, **dict(zip(TABLE_COLUMNS, values))})
        if header is not None:
            sheets[ws.title] = rows
    wb.close()
    return sheets


def normalise_skill(name, catalogue):
    """(skill_name, error) - canonical spelling or 'NEW: <term>'."""
    if name.upper().startswith(NEW_PREFIX):
        term = " ".join(name[len(NEW_PREFIX):].split())
        if not term:
            return None, "NEW: with no term"
        known = catalogue.lookup(term)
        if known:
            return None, f"'{name}' is already in the library as '{known}'"
        return f"{NEW_PREFIX} {term}", None
    canonical = catalogue.by_key.get(skill_key(name))
    if canonical:
        return canonical, None
    hint = catalogue.lookup(name)
    return None, (f"'{name}' is not a canonical name"
                  + (f" (did you mean '{hint}'?)" if hint else " (use a name from the list or NEW: <term>)"))


@dataclass
class Result:
    labels: list = field(default_factory=list)    # rows for the output CSV
    errors: list = field(default_factory=list)    # (jd_ref, excel_row, message)
    warnings: list = field(default_factory=list)  # (jd_ref, excel_row, message)
    labelled: set = field(default_factory=set)    # jd_refs with at least one decision
    unreviewed: dict = field(default_factory=dict)  # A: jd_ref -> rows with keep empty
    merged: int = 0


def collect(sheets, annotator, jd_texts, catalogue, skip_invalid=False):
    """Turn workbook rows into gold label rows. `jd_texts` = {jd_ref: raw_text}."""
    res = Result()
    for ref, rows in sheets.items():
        if ref not in jd_texts:
            res.errors.append((ref, None, "sheet name is not a jd_ref in the sample"))
            continue
        best = {}  # skill key -> label row
        for row in rows:
            if not (row["skill_name"] or row["importance"] or row["evidence_text"]):
                continue  # empty row, or a note on its own
            keep = row["keep"].lower()
            where = (ref, row["excel_row"])
            if keep and keep not in KEEP_YES | KEEP_NO:
                res.errors.append((*where, f"keep must be Y or N, got '{row['keep']}'"))
                continue
            if keep:
                res.labelled.add(ref)
            if keep in KEEP_NO:
                continue
            if annotator == "A" and not keep:
                res.unreviewed[ref] = res.unreviewed.get(ref, 0) + 1
                continue
            if annotator == "B":
                res.labelled.add(ref)

            problems = []
            skill, err = (normalise_skill(row["skill_name"], catalogue) if row["skill_name"]
                          else (None, "skill_name is empty"))
            if err:
                problems.append(err)
            importance = row["importance"].lower()
            if importance not in IMPORTANCE_RANK:
                problems.append(f"importance must be one of {', '.join(IMPORTANCE)}, "
                                f"got '{row['importance']}'")
            if problems:
                res.errors += [(*where, p) for p in problems]
                continue

            evidence = " ".join(row["evidence_text"].split())
            if not evidence:
                res.warnings.append((*where, f"no evidence_text for {skill}"))
            elif not evidence_in_text(evidence, jd_texts[ref]):
                res.warnings.append((*where, f"evidence for {skill} not found in the JD text: "
                                             f"'{evidence[:80]}'"))

            label = {"jd_ref": ref, "skill_name": skill, "importance": importance,
                     "evidence_text": evidence}
            key = skill_key(skill)
            if key in best:
                res.merged += 1
                if IMPORTANCE_RANK[importance] < IMPORTANCE_RANK[best[key]["importance"]]:
                    best[key] = label
            else:
                best[key] = label
        res.labels += best.values()
    if res.errors and not skip_invalid:
        res.labels = []
    res.labels.sort(key=lambda r: r["jd_ref"])  # stable: keeps the sheet order inside a JD
    return res


def role_summary(res, sample, draft_rows, sheets):
    """One dict per role: JDs labelled, rows, NEW rows, draft rows dropped / added by hand.
    The draft comparison only covers JDs that have been labelled."""
    role_of = {r["jd_ref"]: r["role_code"] for r in sample}
    draft = {(d["jd_ref"], skill_key(d["skill_name"])) for d in draft_rows if d["jd_ref"] in res.labelled}
    final = {(r["jd_ref"], skill_key(r["skill_name"])) for r in res.labels}
    out = {}
    for ref in sheets:
        role = role_of.get(ref, "?")
        s = out.setdefault(role, {"role": role, "jds_in_workbook": 0, "jds_labelled": 0, "rows": 0,
                                  "new_rows": 0, "dropped_from_draft": 0, "added_by_hand": 0})
        s["jds_in_workbook"] += 1
        s["jds_labelled"] += ref in res.labelled
    for r in res.labels:
        s = out[role_of[r["jd_ref"]]]
        s["rows"] += 1
        s["new_rows"] += r["skill_name"].startswith(NEW_PREFIX)
        s["added_by_hand"] += (r["jd_ref"], skill_key(r["skill_name"])) not in draft
    for ref, key in draft - final:
        out[role_of[ref]]["dropped_from_draft"] += 1
    return [out[k] for k in sorted(out)]


def print_summary(summary):
    cols = list(summary[0]) if summary else []
    widths = [max(len(c), *(len(str(s[c])) for s in summary)) for c in cols]
    print("  ".join(c.ljust(w) for c, w in zip(cols, widths)))
    totals = {c: sum(s[c] for s in summary) for c in cols[1:]}
    for s in summary + [{"role": "TOTAL", **totals}]:
        print("  ".join(str(s[c]).ljust(w) for c, w in zip(cols, widths)))


def print_issues(title, issues):
    if issues:
        print(f"\n{title} ({len(issues)}):")
        for ref, row, msg in issues:
            print(f"  {ref}" + (f" row {row}" if row else "") + f": {msg}")


def main(argv=None):
    ap = argparse.ArgumentParser(description="Read a filled gold workbook back into a CSV")
    ap.add_argument("workbook", nargs="?", help="default: gold_workbook.xlsx / gold_workbook_B.xlsx")
    ap.add_argument("--annotator", choices=["A", "B"], default="A")
    ap.add_argument("--sample", default=str(SAMPLE_PATH))
    ap.add_argument("--draft", default=str(DRAFT_PATH))
    ap.add_argument("--skills", default=str(SKILLS_PATH))
    ap.add_argument("--out", help="default: gold_labels.csv (A) / gold_labels_B.csv (B)")
    ap.add_argument("--skip-invalid", action="store_true",
                    help="drop rows with errors and write the CSV anyway")
    args = ap.parse_args(argv)

    sample = read_csv(args.sample)
    sheets = read_workbook(args.workbook or WORKBOOK_PATH[args.annotator])
    res = collect(sheets, args.annotator, {r["jd_ref"]: r["raw_text"] for r in sample},
                  Catalogue(load_skills(args.skills)), skip_invalid=args.skip_invalid)

    print_issues("Warnings", res.warnings)
    started = {ref: n for ref, n in res.unreviewed.items() if ref in res.labelled}
    if started:
        print("\nRows with keep empty in JDs already started (mark them Y or N, they are dropped):")
        for ref, n in sorted(started.items()):
            print(f"  {ref}: {n}")
    not_started = sorted(set(sheets) - res.labelled)
    if not_started:
        print(f"\nNot started ({len(not_started)} JDs): {' '.join(not_started)}")
    print_issues("Errors", res.errors)
    print()
    print_summary(role_summary(res, sample, read_csv(args.draft), sheets))
    if res.merged:
        print(f"\n{res.merged} duplicate skill row(s) merged (strongest importance kept)")
    if res.errors and not args.skip_invalid:
        print(f"\n{len(res.errors)} error(s): nothing written. Fix them in the workbook, "
              "or use --skip-invalid to drop those rows.")
        return 1
    out = args.out or LABELS_PATH[args.annotator]
    write_csv(out, res.labels, LABEL_FIELDS)
    print(f"\n{out}: {len(res.labels)} labels from {len(res.labelled)} JDs")
    return 0


if __name__ == "__main__":
    sys.exit(main())
