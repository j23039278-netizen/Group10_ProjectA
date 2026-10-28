"""
SEAGAS — New student intake (how student accounts are created)

1. An advisor uploads a list of new students (Excel .xlsx or .csv) with their
   full name, phone number, personal email, course (Computer Science, Business
   or ICT) and programme (one of the course's programmes).  POST /batches
2. An admin opens the list, lets the system suggest a student number and a
   school email for each student (both editable), then creates the accounts.
   Each student gets a temporary password.  POST /batches/{id}/assign,
   PATCH /rows/{id}, POST /batches/{id}/create-accounts, POST /rows/{id}/reject
3. The student's personal email receives the student ID, school email and
   temporary password (notifications.py; stored in the outbox when no mail
   server is configured).                                       GET /outbox
4. The student signs in with school email + student ID + temporary password and
   must choose a new password first (routers/auth.py).

Personal emails and phone numbers stay in student_intake_rows: they are only
needed to send the account email. Use made-up data when testing (NFR-05).
"""

import csv
import io
import re
import secrets
from typing import List, Literal, Optional
from uuid import UUID

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import Response
from pydantic import BaseModel, Field
from sqlalchemy import text
from sqlalchemy.orm import Session

import audit
import notifications
import programmes
from advisor_scope import courses_of
from database import get_db
from routers.auth import ROLE_EMAIL_DOMAINS, get_current_user, hash_password, password_error

router = APIRouter()

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_ROWS = 1000
STUDENT_DOMAIN = ROLE_EMAIL_DOMAINS["student"]
STUDENT_NUMBER_RE = re.compile(r"^[0-9]{9}$")
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
PHONE_RE = re.compile(r"^\+?[0-9 ()\-]{7,25}$")

# Header words that identify each column (case-insensitive, first match wins)
COLUMN_WORDS = {
    "full_name":      ("full name", "name", "student name"),
    "phone":          ("phone", "mobile", "contact", "tel"),
    "personal_email": ("personal email", "email", "e-mail"),
    "course":         ("course", "faculty"),
    "programme":      ("programme", "program", "major"),
}
REQUIRED_COLUMNS = "Full Name, Phone Number, Personal Email, Course, Programme"

# What the Course column may say (case-insensitive) -> stored code
COURSES = {
    "cs": "CS", "computer science": "CS",
    "business": "BUSINESS",
    "ict": "ICT", "information and communication technology": "ICT",
    "information & communication technology": "ICT",
}
COURSE_NAMES = {"CS": "Computer Science", "BUSINESS": "Business", "ICT": "ICT"}


def parse_course(value):
    """CS / BUSINESS / ICT, or None if the text is not one of the three courses"""
    return COURSES.get(" ".join((value or "").lower().split()))

ROW_COLUMNS = """r.row_id, r.batch_id, r.line_number, r.full_name, r.phone, r.personal_email, r.course, r.programme,
                 r.student_number, r.school_email, r.status, r.user_id, r.note, r.processed_at"""


# ── Permission checks ─────────────────────────────────────────
def require_staff(user):
    if user.role not in ("advisor", "admin"):
        raise HTTPException(status_code=403, detail="Access denied. Advisor or Admin role required.")


def require_admin(user):
    if user.role != "admin":
        raise HTTPException(status_code=403, detail="Access denied. Admin role required.")


def parse_id(value, what="Not found"):
    """Malformed ids are just "not found", like the other ID routes"""
    try:
        return str(UUID(value))
    except ValueError:
        raise HTTPException(status_code=404, detail=what)


def get_batch(db, batch_id, user):
    batch = db.execute(text("SELECT * FROM student_intake_batches WHERE batch_id = :b"),
                       {"b": parse_id(batch_id, "List not found")}).fetchone()
    # Advisors see only the lists they uploaded; admins see all
    if not batch or (user.role != "admin" and str(batch.uploaded_by) != str(user.user_id)):
        raise HTTPException(status_code=404, detail="List not found")
    return batch


# ── Reading the uploaded file ─────────────────────────────────
def read_table(name, data):
    """Rows of cell strings from a .csv or .xlsx file (first row = headers)"""
    lower = (name or "").lower()
    if lower.endswith(".xlsx"):
        if not data.startswith(b"PK"):
            raise HTTPException(status_code=400, detail="This is not a valid .xlsx file.")
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise HTTPException(status_code=400, detail="The Excel file could not be read.")
        ws = wb.worksheets[0]
        rows = [["" if c is None else str(c).strip() for c in row]
                for row in ws.iter_rows(values_only=True, max_row=MAX_ROWS + 2)]
        wb.close()
        return rows
    if lower.endswith(".csv"):
        try:
            content = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = data.decode("latin-1")
        return [[c.strip() for c in row] for row in csv.reader(io.StringIO(content))]
    raise HTTPException(status_code=400, detail="Upload an Excel (.xlsx) or CSV (.csv) file.")


def find_columns(header):
    cols = {}
    names = [h.lower().strip() for h in header]
    for key, words in COLUMN_WORDS.items():
        for word in words:
            hit = next((i for i, h in enumerate(names) if word in h and i not in cols.values()), None)
            if hit is not None:
                cols[key] = hit
                break
    missing = [k for k in ("full_name", "personal_email", "course", "programme") if k not in cols]
    if missing:
        raise HTTPException(
            status_code=400,
            detail=f"The first row must contain the column names: {REQUIRED_COLUMNS}.")
    return cols


def clean_row(cells, cols, course_programmes):
    """(values, None) or (None, error message) for one data row.
    course_programmes: programmes.by_course(db)"""
    cell = lambda k: (cells[cols[k]] if k in cols and cols[k] < len(cells) else "").strip()
    name, email, phone = cell("full_name"), cell("personal_email").lower(), cell("phone")
    if phone.endswith(".0") and phone[:-2].isdigit():   # Excel stores numbers as floats
        phone = phone[:-2]
    if not name:
        return None, "full name is empty"
    if len(name) > 255:
        return None, "full name is too long"
    if not EMAIL_RE.match(email) or len(email) > 255:
        return None, "personal email is not valid"
    if phone and not PHONE_RE.match(phone):
        return None, "phone number is not valid"
    course = parse_course(cell("course"))
    if not course:
        return None, "course must be Computer Science, Business or ICT"
    programme = course_programmes.get(course, {}).get(" ".join(cell("programme").lower().split()))
    if not programme:
        return None, f"programme must be one of the {COURSE_NAMES[course]} programmes"
    return {"full_name": name, "personal_email": email, "phone": phone or None, "course": course,
            "programme": programme}, None


# ── Advisor: Excel template (Course and Programme are drop-down lists) ──
TEMPLATE_ROWS = [
    ("Example Student", "+60 12-345 6789", "example.student@gmail.com", "Computer Science",
     "Bachelor of Computer Science (Software Development)"),
    ("Another Student", "+60 13-222 3333", "another.student@gmail.com", "Business",
     "Bachelor of Business (Marketing)"),
    ("Third Student",   "+60 14-444 5555", "third.student@gmail.com",   "ICT",
     "Bachelor of Information and Communication Technology (Software Systems)"),
]


@router.get("/template")
def download_template(current_user=Depends(get_current_user), db: Session = Depends(get_db)):
    """An .xlsx with the five columns. Course only accepts the courses this advisor may add;
    Programme only the programmes of the course chosen in the same row."""
    require_staff(current_user)
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation

    codes = [c for c in COURSE_NAMES if c in courses_of(current_user)]
    names = [COURSE_NAMES[c] for c in codes]
    wb = Workbook()
    ws = wb.active
    ws.title = "New students"
    ws.append(["Full Name", "Phone Number", "Personal Email", "Course", "Programme"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in TEMPLATE_ROWS:
        if row[3] in names:
            ws.append(row)
    for col, width in zip("ABCDE", (28, 20, 34, 20, 70)):
        ws.column_dimensions[col].width = width

    # A hidden sheet with each course's programmes; one named range per course
    # (Computer_Science, Business, ICT) so the Programme list can follow the Course cell
    lists = wb.create_sheet("Lists")
    lists.sheet_state = "hidden"
    by = programmes.by_course(db)
    for col, (code, course_name) in enumerate(zip(codes, names), start=1):
        letter = chr(64 + col)
        lists.cell(row=1, column=col, value=course_name)
        items = list(by.get(code, {}).values())
        for i, prog in enumerate(items, start=2):
            lists.cell(row=i, column=col, value=prog)
        ref = f"Lists!${letter}$2:${letter}${max(len(items), 1) + 1}"
        wb.defined_names[course_name.replace(" ", "_")] = DefinedName(course_name.replace(" ", "_"), attr_text=ref)

    course_dv = DataValidation(type="list", formula1=f'"{",".join(names)}"', allow_blank=False,
                               showErrorMessage=True, errorTitle="Course",
                               error=f"Choose {' or '.join(names)} from the list.")
    programme_dv = DataValidation(type="list", formula1='INDIRECT(SUBSTITUTE($D2," ","_"))', allow_blank=False,
                                  showErrorMessage=True, errorTitle="Programme",
                                  error="Choose the course first, then one of its programmes from the list.")
    ws.add_data_validation(course_dv)
    ws.add_data_validation(programme_dv)
    course_dv.add("D2:D1001")      # one row per student, up to the 1,000-student limit
    programme_dv.add("E2:E1001")
    ws.freeze_panes = "A2"

    out = io.BytesIO()
    wb.save(out)
    return Response(
        content=out.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="new_students_template.xlsx"'},
    )


# ── Advisor: upload a list ────────────────────────────────────
@router.post("/batches", status_code=201)
def upload_batch(
    request: Request,
    file: UploadFile = File(...),
    current_user=Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Upload new students (Full Name, Phone Number, Personal Email, Course, Programme). Bad rows are
    skipped and listed."""
    require_staff(current_user)
    data = file.file.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="The file is larger than 2 MB.")
    table = [r for r in read_table(file.filename, data) if any(c for c in r)]
    if len(table) < 2:
        raise HTTPException(status_code=400, detail="The file has no student rows.")
    if len(table) - 1 > MAX_ROWS:
        raise HTTPException(status_code=400, detail=f"At most {MAX_ROWS} students per file.")
    cols = find_columns(table[0])

    # Personal emails already on another list (pending or with an account)
    taken = {r[0] for r in db.execute(
        text("SELECT lower(personal_email) FROM student_intake_rows WHERE status <> 'rejected'")).fetchall()}
    rows, skipped, seen = [], [], set()
    course_programmes = programmes.by_course(db)
    for line, cells in enumerate(table[1:], start=2):
        values, error = clean_row(cells, cols, course_programmes)
        if not error and values["course"] not in courses_of(current_user):
            error = f"you can only add {' / '.join(COURSE_NAMES[c] for c in courses_of(current_user))} students"
        if not error and values["personal_email"] in seen:
            error = "same personal email appears twice in this file"
        if not error and values["personal_email"] in taken:
            error = "this personal email is already on an earlier list"
        if error:
            skipped.append({"line": line, "error": error})
            continue
        seen.add(values["personal_email"])
        rows.append({**values, "line": line})
    if not rows:
        raise HTTPException(status_code=400, detail={"message": "No valid student rows.", "skipped": skipped})

    name = re.split(r"[\\/]", file.filename or "students")[-1][:255]
    batch_id = db.execute(
        text("INSERT INTO student_intake_batches (uploaded_by, file_name) VALUES (:u, :n) RETURNING batch_id"),
        {"u": str(current_user.user_id), "n": name}
    ).scalar()
    db.execute(
        text("""
            INSERT INTO student_intake_rows (batch_id, line_number, full_name, phone, personal_email, course, programme)
            VALUES (:b, :line, :full_name, :phone, :personal_email, :course, :programme)
        """),
        [{"b": str(batch_id), **r} for r in rows]
    )
    audit.record(db, "intake.upload", user=current_user, request=request,
                 target_type="intake_batch", target_id=batch_id,
                 details={"file": name, "students": len(rows), "skipped": len(skipped)})
    db.commit()
    return {"batch_id": str(batch_id), "students": len(rows), "skipped": skipped}


@router.get("/batches")
def list_batches(current_user=Depends(get_current_user), db: Session = Depends(get_db)):
    """Uploaded lists, newest first, with how many students are pending / created / rejected"""
    require_staff(current_user)
    rows = db.execute(text("""
        SELECT b.batch_id, b.file_name, b.uploaded_at, u.full_name AS uploaded_by_name,
               COUNT(r.row_id)                                    AS total,
               COUNT(r.row_id) FILTER (WHERE r.status = 'pending')  AS pending,
               COUNT(r.row_id) FILTER (WHERE r.status = 'created')  AS created,
               COUNT(r.row_id) FILTER (WHERE r.status = 'rejected') AS rejected
        FROM student_intake_batches b
        LEFT JOIN users u ON u.user_id = b.uploaded_by
        LEFT JOIN student_intake_rows r ON r.batch_id = b.batch_id
        WHERE (:admin OR b.uploaded_by = CAST(:uid AS uuid))
        GROUP BY b.batch_id, u.full_name
        ORDER BY b.uploaded_at DESC
    """), {"admin": current_user.role == "admin", "uid": str(current_user.user_id)}).fetchall()
    return [dict(r._mapping) for r in rows]


@router.get("/batches/{batch_id}")
def get_batch_rows(batch_id: str, current_user=Depends(get_current_user), db: Session = Depends(get_db)):
    require_staff(current_user)
    batch = get_batch(db, batch_id, current_user)
    rows = db.execute(text(f"""
        SELECT {ROW_COLUMNS} FROM student_intake_rows r
        WHERE r.batch_id = :b ORDER BY r.line_number
    """), {"b": str(batch.batch_id)}).fetchall()
    return {"batch": dict(batch._mapping), "rows": [dict(r._mapping) for r in rows]}


# ── Admin: assign numbers and create accounts ─────────────────
def school_email_for(student_number):
    return f"{student_number}@{STUDENT_DOMAIN}"


def check_assignment(db, row_id, student_number, school_email):
    """Raise 400 if the number or email is malformed or already used"""
    if not STUDENT_NUMBER_RE.match(student_number or ""):
        raise HTTPException(status_code=400, detail="A student ID must be exactly 9 digits, e.g. 105967029.")
    if not (school_email or "").lower().endswith("@" + STUDENT_DOMAIN) or not EMAIL_RE.match(school_email):
        raise HTTPException(status_code=400, detail=f"The school email must end with @{STUDENT_DOMAIN}.")
    if db.execute(text("SELECT 1 FROM users WHERE student_number = :n"), {"n": student_number}).fetchone():
        raise HTTPException(status_code=400, detail=f"Student ID {student_number} already belongs to an account.")
    if db.execute(text("SELECT 1 FROM users WHERE lower(email) = lower(:e)"), {"e": school_email}).fetchone():
        raise HTTPException(status_code=400, detail=f"{school_email} already belongs to an account.")
    clash = db.execute(text("""
        SELECT 1 FROM student_intake_rows
        WHERE status = 'pending' AND row_id <> CAST(:r AS uuid)
          AND (student_number = :n OR lower(school_email) = lower(:e))
    """), {"r": str(row_id), "n": student_number, "e": school_email}).fetchone()
    if clash:
        raise HTTPException(status_code=400, detail="Another new student on a list already has this student ID or email.")


@router.post("/batches/{batch_id}/assign")
def assign_numbers(batch_id: str, current_user=Depends(get_current_user), db: Session = Depends(get_db)):
    """Suggest a student ID and school email for every pending student that has none"""
    require_admin(current_user)
    batch = get_batch(db, batch_id, current_user)
    rows = db.execute(text("""
        SELECT row_id FROM student_intake_rows
        WHERE batch_id = :b AND status = 'pending' AND student_number IS NULL
        ORDER BY line_number
    """), {"b": str(batch.batch_id)}).fetchall()
    for r in rows:
        # next_student_number() skips numbers that accounts already use; the
        # sequence never hands out the same number twice
        n = db.execute(text("SELECT next_student_number('')")).scalar()
        db.execute(text("UPDATE student_intake_rows SET student_number = :n, school_email = :e WHERE row_id = :r"),
                   {"n": n, "e": school_email_for(n), "r": str(r.row_id)})
    db.commit()
    return get_batch_rows(batch_id, current_user, db)


class RowUpdate(BaseModel):
    # Student ID and school email are changed together; course and programme on their own
    student_number: Optional[str] = Field(None, min_length=9, max_length=9)
    school_email: Optional[str] = Field(None, min_length=3, max_length=255)
    course: Optional[Literal["CS", "BUSINESS", "ICT"]] = None
    programme: Optional[str] = Field(None, max_length=150)


@router.patch("/rows/{row_id}")
def update_row(row_id: str, req: RowUpdate, current_user=Depends(get_current_user), db: Session = Depends(get_db)):
    """Change the suggested student ID / school email, the course or the programme of a pending student"""
    require_admin(current_user)
    rid = parse_id(row_id, "Student not found")
    row = db.execute(text("SELECT status, course, programme FROM student_intake_rows WHERE row_id = :r"),
                     {"r": rid}).fetchone()
    if not row:
        raise HTTPException(status_code=404, detail="Student not found")
    if row.status != "pending":
        raise HTTPException(status_code=400, detail="Only students still waiting for an account can be changed.")
    if req.student_number is None and req.school_email is None and req.course is None and req.programme is None:
        raise HTTPException(status_code=400, detail="Nothing to update.")
    if (req.student_number is None) != (req.school_email is None):
        raise HTTPException(status_code=400, detail="Give both the student ID and the school email.")
    if req.student_number is not None:
        number, email = req.student_number.strip(), req.school_email.strip().lower()
        check_assignment(db, rid, number, email)
        db.execute(text("UPDATE student_intake_rows SET student_number = :n, school_email = :e WHERE row_id = :r"),
                   {"n": number, "e": email, "r": rid})
    if req.course is not None or req.programme is not None:
        course = req.course or row.course
        programme = row.programme
        if req.programme is not None:
            programme = programmes.check_programme(db, req.programme, course)
        elif programme and programme not in programmes.by_course(db).get(course, {}).values():
            programme = None        # the old programme belongs to the old course: choose again
        db.execute(text("UPDATE student_intake_rows SET course = :c, programme = :p WHERE row_id = :r"),
                   {"c": course, "p": programme, "r": rid})
    db.commit()
    return dict(db.execute(text(f"SELECT {ROW_COLUMNS} FROM student_intake_rows r WHERE r.row_id = :r"),
                           {"r": rid}).fetchone()._mapping)


def temporary_password(length=12):
    """Random and readable (no 0/O, 1/l/I); meets the password rules in routers/auth.py"""
    upper = "ABCDEFGHJKLMNPQRSTUVWXYZ"
    lower = "abcdefghijkmnopqrstuvwxyz"
    digits = "23456789"
    symbols = "!@#$%&*?"
    pool = upper + lower + digits + symbols
    while True:
        pw = "".join(secrets.choice(pool) for _ in range(length))
        if password_error(pw) is None:
            return pw


class CreateAccounts(BaseModel):
    row_ids: Optional[List[UUID]] = None   # None = every pending student in the list


@router.post("/batches/{batch_id}/create-accounts")
def create_accounts(
    batch_id: str,
    req: CreateAccounts,
    request: Request,
    current_user=Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """Create student accounts and email each student their login details.

    A student without an assigned ID gets a suggested one. A student whose ID or
    email is already taken is skipped and reported; the others still go ahead.
    """
    require_admin(current_user)
    batch = get_batch(db, batch_id, current_user)
    params = {"b": str(batch.batch_id)}
    only = ""
    if req.row_ids is not None:
        only = "AND row_id = ANY(CAST(:ids AS uuid[]))"
        params["ids"] = [str(i) for i in req.row_ids]
    rows = db.execute(text(f"""
        SELECT row_id, full_name, personal_email, student_number, school_email, course, programme
        FROM student_intake_rows
        WHERE batch_id = :b AND status = 'pending' {only}
        ORDER BY line_number
    """), params).fetchall()

    created, failed = [], []
    for r in rows:
        if not r.course:   # lists uploaded before the Course column existed
            failed.append({"row_id": str(r.row_id), "full_name": r.full_name,
                           "error": "no course: choose Computer Science, Business or ICT first"})
            continue
        if not r.programme:   # lists uploaded before the Programme column existed
            failed.append({"row_id": str(r.row_id), "full_name": r.full_name,
                           "error": "no programme: choose the student's programme first"})
            continue
        number, email = r.student_number, r.school_email
        if not number:
            number = db.execute(text("SELECT next_student_number('')")).scalar()
            email = school_email_for(number)
        try:
            check_assignment(db, r.row_id, number, email)
        except HTTPException as e:
            failed.append({"row_id": str(r.row_id), "full_name": r.full_name, "error": e.detail})
            continue
        password = temporary_password()
        user_id = db.execute(text("""
            INSERT INTO users (email, password_hash, role, full_name, student_number, course, must_change_password)
            VALUES (:e, :pw, 'student', :name, :n, :course, TRUE)
            RETURNING user_id
        """), {"e": email, "pw": hash_password(password), "name": r.full_name, "n": number,
               "course": r.course}).scalar()
        programmes.set_programme(db, user_id, r.programme)
        db.execute(text("""
            UPDATE student_intake_rows
            SET status = 'created', user_id = :u, student_number = :n, school_email = :e,
                processed_by = :admin, processed_at = CURRENT_TIMESTAMP
            WHERE row_id = :r
        """), {"u": str(user_id), "n": number, "e": email,
               "admin": str(current_user.user_id), "r": str(r.row_id)})
        subject, body = notifications.welcome_email(r.full_name, number, email, password, COURSE_NAMES[r.course],
                                                    r.programme)
        mail = notifications.send_email(db, r.personal_email, subject, body, related_user_id=user_id)
        created.append({"row_id": str(r.row_id), "full_name": r.full_name, "student_number": number,
                        "school_email": email, "email_status": mail})

    if created:
        audit.record(db, "intake.accounts_create", user=current_user, request=request,
                     target_type="intake_batch", target_id=batch.batch_id,
                     details={"created": len(created), "failed": len(failed)})
    db.commit()
    return {"created": created, "failed": failed, "mail_server": notifications.smtp_configured()}


class RowReject(BaseModel):
    note: str = Field("", max_length=500)


@router.post("/rows/{row_id}/reject")
def reject_row(row_id: str, req: RowReject, request: Request,
               current_user=Depends(get_current_user), db: Session = Depends(get_db)):
    """Do not create an account for this student (duplicate, wrong person ...)"""
    require_admin(current_user)
    rid = parse_id(row_id, "Student not found")
    done = db.execute(text("""
        UPDATE student_intake_rows
        SET status = 'rejected', note = :note, processed_by = :admin, processed_at = CURRENT_TIMESTAMP
        WHERE row_id = :r AND status = 'pending'
        RETURNING batch_id
    """), {"note": req.note.strip() or None, "admin": str(current_user.user_id), "r": rid}).fetchone()
    if not done:
        raise HTTPException(status_code=404, detail="No pending student found.")
    audit.record(db, "intake.reject", user=current_user, request=request,
                 target_type="intake_row", target_id=rid, details={"note": req.note.strip()})
    db.commit()
    return {"message": "Rejected"}


@router.get("/outbox")
def list_outbox(limit: int = 50, current_user=Depends(get_current_user), db: Session = Depends(get_db)):
    """Notification emails, newest first. Without a mail server they are only stored here."""
    require_admin(current_user)
    rows = db.execute(text("""
        SELECT email_id, to_address, subject, body, status, error, created_at
        FROM email_outbox ORDER BY created_at DESC LIMIT :n
    """), {"n": max(1, min(limit, 200))}).fetchall()
    return {"mail_server": notifications.smtp_configured(), "emails": [dict(r._mapping) for r in rows]}
